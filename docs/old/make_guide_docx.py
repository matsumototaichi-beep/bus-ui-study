# -*- coding: utf-8 -*-
"""参加者にわたす案内を Word にする（募集にも当日にも使う）

    python tools/make_guide_docx.py

★UI研究/参加者向け案内.docx が出る。** で挟んだところが太字になる。
"""
import os
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Mm, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUTDIR = os.path.join(os.path.dirname(ROOT), "★UI研究")
OUT = os.path.join(OUTDIR, "参加者向け案内.docx")
FONT = "Yu Gothic"
BLUE = (0x1A, 0x4B, 0x8C)


def jp(run, size, bold=False, color=None):
    run.font.size = Pt(size)
    run.bold = bold
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    if color:
        run.font.color.rgb = RGBColor(*color)


def p(doc, text="", size=11, bold=False, color=None, before=0, after=5):
    par = doc.add_paragraph()
    par.paragraph_format.space_before = Pt(before)
    par.paragraph_format.space_after = Pt(after)
    for i, chunk in enumerate(text.split("**")):
        if chunk:
            jp(par.add_run(chunk), size, bold or i % 2 == 1, color)
    return par


def h1(doc, text, sub):
    p(doc, text, size=18, bold=True, after=2)
    p(doc, sub, size=11, color=(0x55, 0x55, 0x55), after=6)


def h2(doc, text, new_page=False):
    par = p(doc, text, size=13.5, bold=True, color=BLUE, before=7, after=4)
    par.paragraph_format.page_break_before = new_page
    bd = OxmlElement("w:pBdr")
    bot = OxmlElement("w:bottom")
    bot.set(qn("w:val"), "single"); bot.set(qn("w:sz"), "8")
    bot.set(qn("w:space"), "3"); bot.set(qn("w:color"), "1A4B8C")
    bd.append(bot)
    par._p.get_or_add_pPr().append(bd)


# ★2026-10-07: 時刻表のために見出し行つき・列数自由にした。head が無ければ1列目を太字にする（今までと同じ見た目）
def table(doc, body, widths, head=None, gap=True):
    t = doc.add_table(rows=0, cols=len(widths))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    t.autofit = False
    for ri, row in enumerate(([head] if head else []) + list(body)):
        cells = t.add_row().cells
        for ci, val in enumerate(row):
            cells[ci].width = Mm(widths[ci])
            par = cells[ci].paragraphs[0]
            par.paragraph_format.space_before = Pt(2)
            par.paragraph_format.space_after = Pt(2)
            strong = (head and ri == 0) or ci == 0
            for i, chunk in enumerate(str(val).split("**")):
                if chunk:
                    jp(par.add_run(chunk), 11, strong or i % 2 == 1)
    if gap:
        p(doc, "", size=4, after=0)


def steps(doc, items):
    for i, t in enumerate(items, start=1):
        par = doc.add_paragraph()
        par.paragraph_format.space_before = Pt(0)
        par.paragraph_format.space_after = Pt(3)
        par.paragraph_format.left_indent = Mm(6)
        jp(par.add_run("%d.  " % i), 11.5, True, BLUE)
        for j, chunk in enumerate(t.split("**")):
            if chunk:
                jp(par.add_run(chunk), 11.5, j % 2 == 1)


doc = Document()
st = doc.styles["Normal"]
st.font.name = FONT
st.font.size = Pt(11)
st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
st.paragraph_format.line_spacing = 1.0      # ★2026-10-07: A4 2ページに収めるため
for s in doc.sections:
    s.page_width, s.page_height = Mm(210), Mm(297)
    s.top_margin = s.bottom_margin = Mm(12)
    s.left_margin = s.right_margin = Mm(18)

h1(doc, "バスの人数アプリ　実験のご案内", "大和大学　情報学部　松本泰知　卒業研究")

# ★2026-10-07: 入力方式・わからないボタンには触れない。どの回に何が出るかが分かると条件への反応が混ざるため
# ★2026-10-07: 「だいたいでよい」のような、数えにくいときの逃げ道を示す文も書かない（ボタン実験の狙いを打ち消すため）
h2(doc, "どんな研究か")
p(doc, "**バスに乗って、車内の人数をアプリで答えてもらいます。**答えやすい画面を調べる研究です。")
p(doc, "**アプリは、画面の指示に従って操作してください。**", color=BLUE)

h2(doc, "参加の回数")
table(doc, [
    ("UI研究", "**2回**", "発車するたびに、車内の人数をアプリで答える。往復のあとアンケート"),
    ("計数研究", "**1回**", "ゆき：車内の人数を1から数える\nかえり：乗った人・降りた人を数える"),
], [26, 18, 130], head=("", "回数", "すること"))
p(doc, "UI研究は**昼と夕方に1回ずつ**（日にち・順番は自由）。計数研究は日程の最後のほうです。")

h2(doc, "時刻　1回＝往復2本")
table(doc, [
    ("昼", "11:35", "12:03", "12:48", "13:48", "14:32"),
    ("夕方", "15:25", "15:55", "16:40", "17:18", "18:03"),
], [20, 24, 32, 26, 40, 32],
   head=("", "集合", "ゆき 南口発", "桃山台着", "かえり 桃山台発", "南口着"))
p(doc, "集合：ＪＲ吹田駅　**南口**　（北口ではありません）")

h2(doc, "お願い")
table(doc, [
    ("乗るバス", "**上の時刻の便だけ**に乗ってください"),
    ("持ちもの", "充電したスマホ　／　ICカードか小銭　／　2回目は1回目のカード"),
    ("運賃", "**実費をお支払いします。**謝礼はありません"),
    ("やめたいとき", "いつでもやめられます。理由を言う必要はありません"),
], [30, 144])

h2(doc, "記録されること")
# ★2026-10-07: 位置情報は答える15秒前から送ったあと数秒まで記録している。「答えている間」では事実より狭い
p(doc, "答えた混み具合・人数・バス停、答えるのにかかった時間、押したボタン、送った時刻、"
       "**答える前後のスマホの位置情報**。")
p(doc, "名前と電話番号はアプリに入りません。データはIDで扱い、研究以外には使いません。")

h2(doc, "申し込み・連絡先")
table(doc, [
    ("申し込み", "参加できる日と時間帯（昼／夕方）を、松本まで連絡してください"),
    ("連絡先", "松本　泰知　大和大学 情報学部　23610225tm@stu.yamato-u.ac.jp"),
], [30, 144], gap=False)

# ★2026-10-07: 1ページ目＝募集、2ページ目＝当日に見るところ
# ★2026-10-07: 位置情報は報告画面が出た時点で聞かれるので、練習の前に置く。
# かえりは実験者がいないので、のりば・番号・行先・終点まで乗ることを参加者が自分で分かるように書く（北口で降りる・[3][5] に乗るのを防ぐ）
h2(doc, "当日の流れ（UI研究）", new_page=True)
steps(doc, [
    "南口に集合。1回目はカードを受け取る（IDとパスワードが書いてあります）",
    "実験者が見せるQRコードを、スマホのカメラで読む",
    "カードのIDとパスワードでログインする",
    "「**ゆき　吹田駅 → 桃山台駅**」を選び、画面を実験者に見せる",
    "位置情報の利用を聞かれたら「**許可**」",
    "のりばで練習：「**練習**」を押す → のりばで待っている人の数を**2回**答える → "
    "画面の上の「**練習モード**」をタップして終える",
    "実験者が指定したバスに乗る",
    "座ったら「着席」、立っていたら「立席」を押す。変わったら押し直す",
    "**バスが発車するたびに、人数を答える**",
    "桃山台駅で降りたら「**かえりに切り替える**」を押す",
    "かえりは桃山台駅 **2番のりば**から、**13:48／17:18 発**の「**[2] ＪＲ吹田駅（南口）**」に乗る。**終点まで乗る**",
    "かえりのバスでも、発車するたびに人数を答える",
    "南口に着いたら、松本からチャットで届くアンケートに答える。解散",
])

# ★2026-10-07: 「数える人」は「数える係の人」とも読めるので「数に入れる人」にした。
# 「次のバス停に着くまで」は通過するバス停でも区切りになるか曖昧なので「停まるまで」にした
h2(doc, "数え方")
table(doc, [
    ("数に入れる人", "車内にいる人全員。**自分も入れる**"),
    ("入れない人", "**運転士**"),
    ("いつ", "バス停を**発車したら**数えて答え、**次のバス停で停まるまでに送る**"),
], [30, 144])

h2(doc, "人数の答え方")
# ★2026-10-07: 最後はアプリのボタン名（「送信」）にそろえる
steps(doc, [
    "混み具合を6つから選ぶ",
    "車内の人数を入れる",
    "次に停まるバス停を選ぶ",
    "「**送信**」を押す",
])

os.makedirs(OUTDIR, exist_ok=True)
doc.save(OUT)
print("書き出しました: " + OUT)
