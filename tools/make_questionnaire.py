# -*- coding: utf-8 -*-
"""事後アンケート用紙を Word で作る（2026-09-26）

`QUESTIONNAIRE.md` v2.0 の設計をそのまま紙の用紙にする。
提示順のカウンターバランスのため **v1〜v4 の4版**を書き出す。
設問はまったく同じで、S2 と S4 の**提示順だけ**が違う。

    python tools/make_questionnaire.py

  docs/questionnaire/アンケート用紙_v1.docx 〜 _v4.docx が出来る。
  **Word なので、あとから自由に書き換えられる。**

★2026-09-26：**画面の写真は使わない。** 入力方式を実施日ごとの固定に変えたので、
参加者は「その日ずっと使っていた方式」として思い出せる。文章で説明すれば足りる。

Google Forms を作るときも、この用紙をそのまま写せば設問の抜けが出ない。
"""
import os
from docx import Document
from docx.shared import Pt, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUTDIR = os.path.join(ROOT, "docs", "questionnaire")
FONT = "Yu Gothic"

# 版ごとの提示順（QUESTIONNAIRE.md「フォームの版」）
VERSIONS = {
    "v1": {"s2": ["A-1", "A-2"], "s4": ["B-1", "B-2"]},
    "v2": {"s2": ["A-2", "A-1"], "s4": ["B-1", "B-2"]},
    "v3": {"s2": ["A-1", "A-2"], "s4": ["B-2", "B-1"]},
    "v4": {"s2": ["A-2", "A-1"], "s4": ["B-2", "B-1"]},
}
SCREENS = {
    "A-1": ("ステッパー方式",
            "「−5」「−1」「＋1」「＋5」のボタンで、数を増やしたり減らしたりして合わせる"),
    "A-2": ("テンキー方式",
            "数字のキーを押して、人数を直接入力する"),
    "B-1": ("「わからない」ボタン あり", ""),
    "B-2": ("「わからない」ボタン なし", ""),
}

LIKERT = "1  まったくそう思わない ・ 2 ・ 3  どちらとも言えない ・ 4 ・ 5  とてもそう思う"


# ---------------------------------------------------------------- 書式の道具
def setup(doc):
    st = doc.styles["Normal"]
    st.font.name = FONT
    st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    for s in doc.sections:
        s.page_width, s.page_height = Mm(210), Mm(297)
        s.top_margin = s.bottom_margin = Mm(18)
        s.left_margin = s.right_margin = Mm(18)


def p(doc, text="", size=10.5, bold=False, color=None, space_before=0, space_after=3,
      align=None, indent=0):
    par = doc.add_paragraph()
    par.paragraph_format.space_before = Pt(space_before)
    par.paragraph_format.space_after = Pt(space_after)
    if indent:
        par.paragraph_format.left_indent = Mm(indent)
    if align:
        par.alignment = align
    run = par.add_run(text)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    if color:
        run.font.color.rgb = RGBColor(*color)
    return par


def section(doc, title, note=None):
    p(doc, title, size=13, bold=True, space_before=10, space_after=2)
    if note:
        p(doc, note, size=8.5, color=(0x55, 0x55, 0x55), space_after=5)


def question(doc, no, text, bold_text=False):
    p(doc, "%s  %s" % (no, text), bold=bold_text, space_before=4, space_after=1)


def choices(doc, items, cols=None):
    p(doc, "　" + "　　".join("□ " + c for c in items), size=10, space_after=3)


def likert(doc):
    p(doc, "　" + LIKERT, size=9, color=(0x33, 0x33, 0x33), space_after=3)


def lines(doc, n=2):
    for _ in range(n):
        p(doc, "　" + "＿" * 44, size=10, color=(0x99, 0x99, 0x99), space_after=5)


def method_box(doc, key):
    """入力方式を文章で示す枠。写真は使わない（2026-09-26）。"""
    label, desc = SCREENS[key]
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    cell = t.cell(0, 0)
    cell.width = Mm(174)
    cp = cell.paragraphs[0]
    r1 = cp.add_run("【%s】　" % label)
    r1.bold = True
    r1.font.size = Pt(12)
    r2 = cp.add_run(desc)
    r2.font.size = Pt(10)
    for r in (r1, r2):
        r.font.name = FONT
        r._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    p(doc, "", size=4, space_after=2)
    return label


def likert_block(doc, label):
    p(doc, "【%s】について、それぞれ当てはまる数字に○をつけてください。" % label,
      size=9.5, space_before=2, space_after=2)
    for no, text in [("2-1", "この方式は押しやすかった"),
                     ("2-2", "この方式は素早く答えられた"),
                     ("2-3", "この方式では自分が思った通りの人数を入力できた"),
                     ("2-4", "この方式は頭を使う・疲れると感じた")]:
        p(doc, "　%s  %s" % (no, text), size=10, space_after=0)
        p(doc, "　　　1 ・ 2 ・ 3 ・ 4 ・ 5", size=10, space_after=3)


# ---------------------------------------------------------------- 本体
def build(version):
    v = VERSIONS[version]
    doc = Document()
    setup(doc)

    p(doc, "バス車内の人数を答えるアプリについて（%s）" % version,
      size=15, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    p(doc, "ご協力ありがとうございました。すべての乗車を終えたあとにお答えください。所要 10〜12分です。",
      size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER, color=(0x55, 0x55, 0x55), space_after=8)

    # ---- S0
    section(doc, "S0. 基本情報")
    question(doc, "0-1", "配布カードに書かれたログインID（p01 など）　【必須】", bold_text=True)
    lines(doc, 1)
    question(doc, "0-2", "参加した日　【必須】")
    p(doc, "　　1日目：＿＿月＿＿日　　　2日目：＿＿月＿＿日", size=10, space_after=4)
    question(doc, "0-3", "この用紙の版（配布カードに書かれています）　【必須】")
    choices(doc, ["v1", "v2", "v3", "v4"])
    question(doc, "0-4", "1本目・2本目に乗ったバスの発車時刻（覚えている範囲で。空欄でも構いません）")
    p(doc, "　　1本目：＿＿：＿＿　　　2本目：＿＿：＿＿", size=10, space_after=4)

    # ---- S1
    section(doc, "S1. 普段のことについて")
    question(doc, "1-1", "年代")
    choices(doc, ["10代", "20代", "30代", "40代以上"])
    question(doc, "1-2", "普段バスを利用する頻度")
    choices(doc, ["ほぼ毎日", "週に数回", "月に数回", "ほとんど乗らない"])
    question(doc, "1-3", "今回の路線に乗ったことがありましたか")
    choices(doc, ["よく乗る", "数回ある", "初めて"])
    question(doc, "1-4", "実験で使った端末")
    choices(doc, ["iPhone", "Android", "その他"])
    question(doc, "1-5", "利き手")
    choices(doc, ["右", "左"])
    question(doc, "1-6", "実験中、スマホは主にどちらの手で操作しましたか")
    choices(doc, ["片手（親指）", "両手", "その時による"])

    # ---- S2
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    section(doc, "S2. 2つの入力方式について",
            "2日間で、人数の入れ方が日によって違いました。次の2つです。")
    first, second = v["s2"]
    question(doc, "2-0", "1日目に使ったのはどちらの方式でしたか", bold_text=True)
    choices(doc, [SCREENS["A-1"][0], SCREENS["A-2"][0], "覚えていない"])
    p(doc, "※ 思い出せる範囲で構いません。分からなければ「覚えていない」を選んでください。",
      size=8.5, color=(0x55, 0x55, 0x55), space_after=6)
    for key in (first, second):
        label = method_box(doc, key)
        likert_block(doc, label)

    lbl1, lbl2 = SCREENS[first][0], SCREENS[second][0]
    question(doc, "2-5", "2つのうち、答えやすかったのはどちらですか", bold_text=True)
    choices(doc, [lbl1, lbl2, "どちらとも言えない"])
    question(doc, "2-6", "それはなぜですか")
    lines(doc, 2)
    question(doc, "2-7", "揺れている車内では、どちらが押し間違えにくかったですか", bold_text=True)
    choices(doc, [lbl1, lbl2, "どちらとも言えない"])
    question(doc, "2-8", "それはなぜですか")
    lines(doc, 2)

    # ---- S3
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    section(doc, "S3. 人数を答えるときのことについて",
            "正解・不正解を見るものではありません。思ったとおりにお答えください。")
    question(doc, "3-1", "人数を答えるとき、実際に一人ずつ数えていましたか")
    choices(doc, ["毎回数えた", "だいたい数えた", "目分量が多かった", "ほとんど目分量"])
    question(doc, "3-2", "数えるのが難しいと感じたのはどんなときですか（いくつでも）")
    choices(doc, ["混んでいた", "自分が立っていた", "揺れた", "降車が近かった", "暗かった", "その他"])
    question(doc, "3-3", "数えきれていないのに、それらしい数を入れて送ったことがありますか",
             bold_text=True)
    choices(doc, ["何度もある", "数回ある", "ない", "覚えていない"])
    question(doc, "3-4", "（ある方）そうしたのはどんなときですか")
    lines(doc, 2)
    question(doc, "3-5", "人数を入れずに空欄のまま送ったことがありますか。あればその理由も")
    lines(doc, 2)
    question(doc, "3-6", "混雑しているときと空いているときで、答えやすさは違いましたか")
    p(doc, "　　1 ・ 2 ・ 3 ・ 4 ・ 5　（1＝違わない … 5＝大きく違った）", size=10, space_after=2)
    lines(doc, 2)
    question(doc, "3-7", "実験中、座っていた割合はどのくらいですか")
    choices(doc, ["ほぼ座席", "半々", "ほぼ立席"])

    # ---- 壁
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    p(doc, "ここでいったん止めてください。", size=14, bold=True,
      align=WD_ALIGN_PARAGRAPH.CENTER, color=(0xC0, 0x00, 0x00), space_before=40, space_after=6)
    p(doc, "この先は、ここまでを回収してからお渡しします。",
      size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER, color=(0x55, 0x55, 0x55))
    p(doc, "（実験者へ：S4 は S3 を回収してから渡すこと。先に見せると S3 の回答が"
           "「ボタンのことを意識した状態」になり、汚染されます。"
           "Word の印刷でページ範囲を分けて刷ってください。）",
      size=8.5, align=WD_ALIGN_PARAGRAPH.CENTER, color=(0x99, 0x99, 0x99), space_before=6)

    # ---- S4
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    section(doc, "S4. 「わからない」ボタンについて",
            "同じ日の1本目と2本目で、画面に1か所だけ違いがありました。")
    question(doc, "4-1", "同じ日の1本目と2本目で、画面に違いがあったことに気づきましたか")
    choices(doc, ["はっきり気づいた", "なんとなく気づいた", "気づかなかった"])
    p(doc, "違いはここです。人数を入れる欄のすぐ下に――", size=9.5, space_before=4, space_after=3)
    for key in v["s4"]:
        label = SCREENS[key][0]
        t = doc.add_table(rows=1, cols=1)
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.LEFT
        cp = t.cell(0, 0).paragraphs[0]
        body = ("「わからない」というボタンが置かれていた便"
                if key == "B-1" else "何も置かれていなかった便")
        r1 = cp.add_run("【%s】　" % label); r1.bold = True; r1.font.size = Pt(11)
        r2 = cp.add_run(body); r2.font.size = Pt(10)
        for r in (r1, r2):
            r.font.name = FONT
            r._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        p(doc, "", size=4, space_after=2)
    question(doc, "4-2", "「わからない」ボタンがあったほうの乗車を覚えていますか")
    choices(doc, ["1本目", "2本目", "覚えていない"])
    question(doc, "4-3", "「わからない」ボタンがあったとき、数えきれないときはどうしていましたか")
    choices(doc, ["ボタンを押した", "それらしい数を入れた", "空欄で送った", "覚えていない"])
    question(doc, "4-4", "「わからない」ボタンがなかったとき、数えきれないときはどうしていましたか",
             bold_text=True)
    choices(doc, ["それらしい数を入れた", "空欄で送った", "クラスだけ選んで送った", "覚えていない"])
    question(doc, "4-5", "どちらのほうが正直に答えられたと感じますか")
    _b = ["ボタンがあった便", "ボタンがなかった便"]
    if v["s4"][0] == "B-2":
        _b.reverse()
    choices(doc, _b + ["変わらない"])
    question(doc, "4-6", "そう思う理由")
    lines(doc, 2)
    question(doc, "4-7", "「わからない」と答えるのに、抵抗や後ろめたさはありましたか")
    p(doc, "　　1 ・ 2 ・ 3 ・ 4 ・ 5　（1＝まったくなかった … 5＝とてもあった）", size=10, space_after=2)
    lines(doc, 2)

    # ---- S5
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    section(doc, "S5. アプリ全体の使いやすさ",
            "それぞれ当てはまる数字に○をつけてください。1＝まったくそう思わない … 5＝とてもそう思う")
    sus = ["このアプリを頻繁に使いたいと思う",
           "このアプリは不必要に複雑だと思う",
           "このアプリは簡単に使えると思う",
           "このアプリを使うには、詳しい人のサポートが必要だと思う",
           "このアプリのさまざまな機能は、うまくまとまっていると思う",
           "このアプリには一貫性のないところが多いと思う",
           "たいていの人はこのアプリの使い方をすぐに覚えられると思う",
           "このアプリはとても使いにくいと思う",
           "このアプリを自信をもって使えた",
           "このアプリを使い始める前に、いろいろ覚える必要があった"]
    for i, text in enumerate(sus, start=1):
        p(doc, "5-%d  %s" % (i, text), size=10, space_before=3, space_after=0)
        p(doc, "　　　1 ・ 2 ・ 3 ・ 4 ・ 5", size=10, space_after=2)

    # ---- S6
    section(doc, "S6. 自由にお書きください")
    for no, text in [("6-1", "人数を答えるとき、いちばん面倒だと感じたことは何ですか"),
                     ("6-2", "「次に停まるバス停」を選ぶとき、困ったことはありましたか"),
                     ("6-3", "このアプリをこう変えたら答えやすくなる、という案があれば"),
                     ("6-4", "実験全体について、気づいたことがあれば")]:
        question(doc, no, text)
        lines(doc, 2)

    p(doc, "ご協力ありがとうございました。", size=11, bold=True,
      align=WD_ALIGN_PARAGRAPH.CENTER, space_before=10)
    return doc


def main():
    os.makedirs(OUTDIR, exist_ok=True)
    for version in VERSIONS:
        path = os.path.join(OUTDIR, "アンケート用紙_%s.docx" % version)
        build(version).save(path)
        print("書き出しました: " + path)
    print("\nWord で開いて自由に編集できます。画面の写真は使いません（2026-09-26 変更）。")
    print("S4 は S3 を回収してから渡します。印刷はページ範囲を分けてください。")


if __name__ == "__main__":
    main()
