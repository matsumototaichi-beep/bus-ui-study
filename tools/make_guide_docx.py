# -*- coding: utf-8 -*-
"""参加者にわたす案内を Word にする

    python tools/make_guide_docx.py

★UI研究/参加者向け案内.docx が出る。** で挟んだところが太字になる。
"""
import os
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Mm, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUTDIR = os.path.join(os.path.dirname(ROOT), "★UI研究")
OUT = os.path.join(OUTDIR, "参加者向け案内.docx")
FONT = "Yu Gothic"
BLUE = (0x1A, 0x4B, 0x8C)


def jp(run, size, bold=False, color=None):
    run.font.size = Pt(size)
    run.bold = bold
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    if color:
        run.font.color.rgb = RGBColor(*color)


def p(doc, text="", size=11, bold=False, color=None, before=0, after=5):
    par = doc.add_paragraph()
    par.paragraph_format.space_before = Pt(before)
    par.paragraph_format.space_after = Pt(after)
    for i, chunk in enumerate(text.split("**")):
        if chunk:
            jp(par.add_run(chunk), size, bold or i % 2 == 1, color)
    return par


def h1(doc, text, sub):
    p(doc, text, size=20, bold=True, after=2)
    p(doc, sub, size=11, color=(0x55, 0x55, 0x55), after=14)


def h2(doc, text):
    par = p(doc, text, size=14, bold=True, color=BLUE, before=14, after=5)
    bd = OxmlElement("w:pBdr")
    bot = OxmlElement("w:bottom")
    bot.set(qn("w:val"), "single"); bot.set(qn("w:sz"), "8")
    bot.set(qn("w:space"), "3"); bot.set(qn("w:color"), "1A4B8C")
    bd.append(bot)
    par._p.get_or_add_pPr().append(bd)


def table(doc, body, widths):
    t = doc.add_table(rows=0, cols=2)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    t.autofit = False
    for k, v in body:
        cells = t.add_row().cells
        for ci, val in enumerate((k, v)):
            cells[ci].width = Mm(widths[ci])
            par = cells[ci].paragraphs[0]
            par.paragraph_format.space_before = Pt(3)
            par.paragraph_format.space_after = Pt(3)
            for i, chunk in enumerate(str(val).split("**")):
                if chunk:
                    jp(par.add_run(chunk), 11, ci == 0 or i % 2 == 1)
    p(doc, "", size=4, after=0)


def steps(doc, items):
    for i, t in enumerate(items, start=1):
        par = doc.add_paragraph()
        par.paragraph_format.space_before = Pt(0)
        par.paragraph_format.space_after = Pt(5)
        par.paragraph_format.left_indent = Mm(6)
        jp(par.add_run("%d.  " % i), 11.5, True, BLUE)
        for j, chunk in enumerate(t.split("**")):
            if chunk:
                jp(par.add_run(chunk), 11.5, j % 2 == 1)


doc = Document()
st = doc.styles["Normal"]
st.font.name = FONT
st.font.size = Pt(11)
st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
for s in doc.sections:
    s.page_width, s.page_height = Mm(210), Mm(297)
    s.top_margin = s.bottom_margin = Mm(18)
    s.left_margin = s.right_margin = Mm(18)

h1(doc, "バスの人数アプリ　実験のご案内", "大和大学　情報学部　松本泰知　卒業研究")

h2(doc, "やること")
p(doc, "バスに乗って、**停まるたびに車内の人数をスマホで答える。**それだけです。")

h2(doc, "1回の実験は、往復2本です")
table(doc, [
    ("集合", "ＪＲ吹田駅　**南口**　（北口ではありません）"),
    ("ゆき", "ＪＲ吹田駅（南口） → 桃山台駅　　約45分"),
    ("待ち時間", "桃山台駅で休憩"),
    ("かえり", "桃山台駅 → ＪＲ吹田駅（南口）　　約45分"),
    ("さいごに", "アンケート　5分ほど"),
], [30, 144])

h2(doc, "当日の流れ")
steps(doc, [
    "南口に集合。カードを受け取る（IDとパスワードが書いてあります）",
    "実験者が出すポスターのQRコードを、スマホのカメラで読む",
    "カードのIDとパスワードでログインする",
    "位置情報の利用を聞かれたら「許可」を押す",
    "のりばで2回だけ練習する",
    "バスに乗る。**停まるたびに人数を答える**",
    "桃山台駅で降りる。休憩",
    "かえりのバスでも同じことをする",
    "南口で解散。アンケートに答える",
])

h2(doc, "人数の答え方")
steps(doc, [
    "混み具合を6つから選ぶ",
    "いま車内にいる人数を入れる",
    "次に停まるバス停を選ぶ",
    "送る",
])

h2(doc, "お願い")
table(doc, [
    ("乗るバス", "**実験者が指定した1本にだけ乗ってください。**"
                 "同じ行先のバスが続けて来ます。実験者が見送ります"),
    ("持ちもの", "充電したスマホ　／　ICカードか小銭"),
    ("運賃", "**実費をお支払いします。**謝礼はありません"),
    ("やめたいとき", "いつでもやめられます。理由を言う必要はありません"),
], [30, 144])

h2(doc, "記録されること")
p(doc, "答えた混み具合と人数、答えるのにかかった時間、押したボタン、送った時刻、"
       "**答えている間のスマホの位置情報**。")
p(doc, "名前と電話番号はアプリに入りません。データはIDで扱い、研究以外には使いません。")

h2(doc, "連絡先")
table(doc, [("松本　泰知", "大和大学 情報学部　23610225tm@stu.yamato-u.ac.jp")], [30, 144])

os.makedirs(OUTDIR, exist_ok=True)
doc.save(OUT)
print("書き出しました: " + OUT)
