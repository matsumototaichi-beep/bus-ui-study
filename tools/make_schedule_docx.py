# -*- coding: utf-8 -*-
"""実施割当と日程を Word（.docx）に出力する（2026-10-07）

docs/実施割当と日程_2026-10-07.pdf と同じ中身。
PDFは文字を直せないので、本文を書き換えたいとき用にこちらを使う。

    python tools/make_schedule_docx.py

日付を入れる作業そのものは schedule.xlsx の「日程」シートのほうが楽
（回の番号を入れると組・入力方式が自動で入る）。この .docx は印刷して配ったり、
文面を直して教授や計数役に渡したりするためのもの。

★本文中の ** で挟んだところが太字になる。
"""
import os
import datetime
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Mm, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "docs", "実施割当と日程_2026-10-07.docx")
FONT = "Yu Gothic"
WEEK = "月火水木金土日"
DATES = [(10, 13), (10, 14), (10, 15), (10, 16), (10, 19), (10, 20), (10, 23),
         (11, 2), (11, 4), (11, 5), (11, 6), (11, 9), (11, 10), (11, 11), (11, 12), (11, 13)]
GRAY = (0x44, 0x44, 0x44)
RED = (0xAA, 0x00, 0x00)


def jp(run, size, bold=False, color=None):
    run.font.size = Pt(size)
    run.bold = bold
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    if color:
        run.font.color.rgb = RGBColor(*color)
    return run


def p(doc, text="", size=9.5, bold=False, color=None, before=0, after=4):
    par = doc.add_paragraph()
    par.paragraph_format.space_before = Pt(before)
    par.paragraph_format.space_after = Pt(after)
    for i, chunk in enumerate(text.split("**")):
        if chunk:
            jp(par.add_run(chunk), size, bold or i % 2 == 1, color)
    return par


def h1(doc, text):
    p(doc, text, size=15, bold=True, after=1)


def h2(doc, text):
    par = p(doc, text, size=11.5, bold=True, before=11, after=3)
    bd = OxmlElement("w:pBdr")
    bot = OxmlElement("w:bottom")
    bot.set(qn("w:val"), "single")
    bot.set(qn("w:sz"), "8")
    bot.set(qn("w:space"), "2")
    bot.set(qn("w:color"), "444444")
    bd.append(bot)
    par._p.get_or_add_pPr().append(bd)


def table(doc, head, body, widths=None, header_col=False):
    t = doc.add_table(rows=0, cols=len(head or body[0]))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    t.autofit = False
    for ri, row in enumerate(([head] if head else []) + body):
        cells = t.add_row().cells
        for ci, val in enumerate(row):
            cell = cells[ci]
            if widths:
                cell.width = Mm(widths[ci])
            par = cell.paragraphs[0]
            par.paragraph_format.space_before = Pt(1)
            par.paragraph_format.space_after = Pt(1)
            is_head = (head and ri == 0) or (header_col and ci == 0)
            for i, chunk in enumerate(str(val).split("**")):
                if chunk:
                    jp(par.add_run(chunk), 8.5, is_head or i % 2 == 1)
    return t


doc = Document()
st = doc.styles["Normal"]
st.font.name = FONT
st.font.size = Pt(9.5)
st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
for s in doc.sections:
    s.page_width, s.page_height = Mm(210), Mm(297)
    s.top_margin = s.bottom_margin = Mm(16)
    s.left_margin = s.right_margin = Mm(16)

h1(doc, "実施割当と日程")
p(doc, "2026年10月7日　／　松本　　＊日付が決まったら 3. に書き込んでください。",
  size=8.5, color=GRAY, after=8)

# ---- 1 -------------------------------------------------------------------
h2(doc, "1. 使う便（どの日も同じ）")
table(doc,
      ["枠", "集合", "1本目", "桃山台着", "2本目", "南口着・解散", "拘束"],
      [["**昼**（空いている）", "11:35", "南口 **12:03**発", "12:48",
        "桃山台 **13:48**発", "14:32", "約3時間"],
       ["**夕方**（混んでいる）", "15:25", "南口 **15:55**発", "16:40",
        "桃山台 **17:18**発", "18:03", "約2時間40分"]],
      widths=[30, 16, 27, 19, 30, 25, 25])
p(doc, "すべて阪急バス 吹田市内線2系統。集合・解散はどちらもＪＲ吹田駅（南口）。"
       "4便とも始発便なので、計数役は車内0人から積み上げられます。",
  size=8.3, color=GRAY, before=3)

# ---- 2 -------------------------------------------------------------------
h2(doc, "2. 実施割当（全10日）")
p(doc, "日付以外はすべて決まっています。本実験が8日、計数方法の実験が2日です。")
body = []
n = 0
for team in (1, 2, 3, 4):
    gun = 1 if team % 2 == 1 else 2
    noon = ("ステッパー", "A", "B") if gun == 1 else ("テンキー", "D", "C")
    eve = ("テンキー", "D", "C") if gun == 1 else ("ステッパー", "A", "B")
    who = "p%02d〜p%02d" % (team * 3 - 2, team * 3)
    for waku, (hoshiki, p1, p2) in (("昼（空いている）", noon), ("夕方（混んでいる）", eve)):
        n += 1
        body.append(["**%d**" % n, waku, "組%d" % team, "群%d" % gun,
                     hoshiki, "**%s**" % p1, "**%s**" % p2, who, ""])
for i in (9, 10):
    body.append(["**%d**" % i, "昼＋夕方（4便）", "―", "―", "―", "―", "―",
                 "計数役3名のみ", ""])
table(doc, ["回", "枠", "組", "群", "入力方式", "1本目", "2本目", "参加者", "日付"], body,
      widths=[9, 30, 13, 13, 21, 14, 14, 23, 25])
p(doc, "1本目・2本目の欄はQRポスターの記号です。"
       "**A＝ステッパー＋ボタンあり　B＝ステッパー＋ボタンなし　"
       "C＝テンキー＋ボタンあり　D＝テンキー＋ボタンなし**", size=8.3, color=GRAY, before=3)
p(doc, "**同じ組の2行（昼と夕方）は、同じ3名が別の日に来ます。**間が空いても構いません。"
       "**組1・組3は昼がステッパー、組2・組4は昼がテンキー**です。揃えてしまうと、"
       "入力方式の差なのか混雑の差なのか分からなくなります。"
       "これでどの参加者も4乗車で「入力方式2通り × ボタン2通り」を1回ずつ経験します。",
  size=8.3, color=GRAY)
p(doc, "回9・10（計数方法の実験）は参加者を呼びません。計数役3名だけで、"
       "1日に昼と夕方の4便を回します。足し引き2名＋数え直し1名で、乗降計測アプリを使います。",
  size=8.3, color=GRAY)

# ---- 3 -------------------------------------------------------------------
h2(doc, "3. 日付ごとの記入表")
p(doc, "使える平日は16日です（10/12 スポーツの日、11/3 文化の日を除く）。"
       "昼枠・夕方枠に、上の「回」の番号を書き込んでください。")
rows = []
for m, d in DATES:
    dt = datetime.date(2026, m, d)
    rows.append(["**%d/%d**" % (m, d), WEEK[dt.weekday()], "", "", ""])
table(doc, ["日付", "曜", "昼枠（11:35集合）に入れる回", "夕方枠（15:25集合）に入れる回", "備考"],
      rows, widths=[20, 11, 55, 55, 37])
p(doc, "本実験8日＋計数方法2日＝**10日**。**予備が6日**残ります。"
       "雨・欠席・遅延はこれで吸収できます。", size=8.3, color=GRAY, before=3)
p(doc, "**土日は使えません。**10/24・25、11/14・15 は土日で、"
       "桃山台17時台が土休日ダイヤでは 北10/20/北40/50 となり **17:18 が存在しない**ためです。"
       "便を選び直すと平日のデータと混ぜられなくなるので、予備日からも外しています。",
  size=8.3, color=GRAY)

# ---- 4 -------------------------------------------------------------------
h2(doc, "4. 当日の流れ")
table(doc, ["時刻", "昼の日", "時刻", "夕方の日"],
      [["11:35", "ＪＲ吹田駅（南口）集合。説明・同意取得・カード配布・ログイン確認",
        "15:25", "同左（2日目の方は説明を短縮）"],
       ["11:50", "のりば2へ移動。のりばで練習2回", "15:42", "のりば2へ移動。のりばで練習2回"],
       ["**12:03**", "**1本目 発車**（松本は離脱）", "**15:55**", "**1本目 発車**（松本は離脱）"],
       ["12:48", "桃山台駅着。折り返し待ち60分", "16:40", "桃山台駅着。折り返し待ち38分"],
       ["**13:48**", "**2本目 発車**", "**17:18**", "**2本目 発車**"],
       ["14:32", "ＪＲ吹田駅（南口）着・解散。アンケート",
        "18:03", "ＪＲ吹田駅（南口）着・解散。アンケート"]],
      widths=[16, 73, 16, 73])
p(doc, "アンケートは解散直後に①前半、送信を確認してから①後半を渡します。"
       "2日目の方は①後半の末尾に追加の質問があります。", size=8.3, color=GRAY, before=3)

# ---- 5 -------------------------------------------------------------------
h2(doc, "5. 当日いちばん事故りやすいところ")
p(doc, "**南口のりば2からは、[2] 桃山台駅ゆき と [3] 桃山台駅ゆき の両方が出ます。**"
       "行先表示はどちらも「桃山台駅」で、**行先番号の数字でしか見分けられません。**"
       "使うのは [2]（12:03／15:55発）。乗ってはいけないのは [3]（毎時32分発。経路が違います）。",
  size=8.5, color=RED)
p(doc, "**桃山台では「[2] ＪＲ吹田駅（南口）ゆき」に乗り、終点の南口まで乗ります。**"
       "「ＪＲ吹田駅（北口）」と表示して発車するのは [5] という別系統です。",
  size=8.5, color=RED)
p(doc, "参加者には「実験者が指定した1本にだけ乗る。来た順に乗らない」と説明し、"
       "**発車時刻を口に出して指定し、松本が見送るまでその場を離れません。**"
       "9月29日に、この注意書きを書いた松本自身が桃山台で [3] に乗っています。"
       "北口を経由すると思い込んでいたためで、行先表示の見落としではありませんでした。"
       "「行先を確認してください」では防げないので、見送りを省略しません。",
  size=8.3, color=GRAY)

# ---- 6 -------------------------------------------------------------------
h2(doc, "6. 持っていくもの")
table(doc, None,
      [["配布カード", "参加者番号・組・群・アンケートの版・ID・パスワード。番号順に3枚ずつ渡す"],
       ["QRポスター", "**A・B・C・D の4枚。**その日に使うのは2枚（2. の1本目・2本目の欄）"],
       ["記録用紙", "22停留所版。1日4枚（2便 × 乗車係・降車係）"],
       ["対応表", "accounts.csv。参加者がパスワードを無くしたとき用。人目に触れない場所に"],
       ["この用紙", "その日の回の行に印を付けておく"]],
      widths=[28, 150], header_col=True)

doc.save(OUT)
print("書き出しました: " + OUT)
