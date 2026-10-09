# -*- coding: utf-8 -*-
"""当日、参加者ひとりずつに配る用紙（A4 両面1枚）を Word と PDF にする

    python tools/make_handouts.py                  本番。アンケートの URL がそろった用紙だけ書き出す
    python tools/make_handouts.py --draft <フォルダ>  下書き。URL がないところは仮の QR にして、3枚とも <フォルダ> へ
        <フォルダ> はこのリポジトリの外にする（URL がそろったアンケートは本物の URL が入るため。中なら止まる）

本番で書き出すもの
  ../★UI研究/配布用紙_UI研究_1回目.docx / .pdf
  ../★UI研究/配布用紙_UI研究_2回目.docx / .pdf
  ../★乗降人数計算/配布用紙_計数研究.docx / .pdf

表＝時間割・お願い・当日の流れ・アプリのQR
裏＝桃山台駅ですること・数え方・答え方・アンケートのQR
   UI研究は アンケートが2つ（桃山台駅で ゆき、南口で かえり）。計数研究は南口で1つ

ログインの ID とパスワードは載せない（今までどおり配布カードを別に渡す）。
アンケートの URL はこのリポジトリに書かない（public のため。知らない人の回答が混ざる）。
実行のたびに アンケートURL.xlsx（シート「アンケートURL」）の「回答用URL（配布用紙のQR）」
または「回答用URL（参加者に送る）」の列と「問数」の列から読む。「使わなくなったもの」より下は読まない。
PDF は Word で書き出す（Word が要る）。一時フォルダで作り、3枚とも2ページちょうどのときだけ
書き出し先へ写す（2ページでなければ、書き出し先の前の版に触らずに止まる）。
** で挟んだところが太字になる。

★2026-10-07: 当日の説明を減らし、桃山台駅に着いてから迷わないように作った。
  文言は参加者向け案内（★UI研究/参加者向け案内_Ui研究.docx）の言い回しに合わせている。
  UI研究の用紙には条件（入力方式・「わからない」ボタン）が分かることを書かない。
  「だいたいでよい」のような、数えにくいときの逃げ道を示す文も書かない。
★2026-10-08: 見直しの指摘で直した。桃山台駅の手順に「着席／立席」（3種類）、
  「よろしいですか？」→「OK」（UI研究）、桃山台駅で乗ってきた人から数えること（計数研究）を足した。
  乗らないバスに「同じ [2] 南口行きでも、ほかの時刻」を分けて書いた。位置情報は「記録します」と書いた。
  「よろしいですか？」「計数研究　かえり」が行の途中で分かれないように \n で改行した。
★2026-10-08(2): UI研究のアンケートを区間ごとに分けた（1回の往復で、桃山台駅と南口の2回答える）。
  裏の「桃山台駅ですること」に ゆき のアンケートの QR と、答えたら人数アプリに戻ることを足し、
  南口の QR は「アンケート　かえり」にした。
  --draft を足した（フォームを作る前に、仮の QR で形を確かめる）。
  文は yomiyasu・natural-japanese の2つのスキルで見直して減らした（意味と手順の順番は変えていない）。
  同じことを2回書いていたところを1回にした（表の流れ10の「かえりに切り替える」は裏の1に任せる。
  裏の「困ったときは…」は表のお願いと同じなので消した）。
  位置情報を聞かれずに拒否されたときの直し方は、アプリの帯と直し方の画面が伝えるので、用紙には書かない。
★2026-10-08(3): 見直しの指摘で直した。--draft にリポジトリの中のフォルダを渡したら止まる。
  用紙は一時フォルダで作り、2ページだと確かめてから書き出し先へ写す。位置情報の2文目の「位置情報から」を
  参加者向け案内に合わせて戻した。「…に答える（6問・1分ほど）。」と句点を打ち、
  「昼は約60分、夕方は約38分」「発車までに2番のりばへ」と、語の間の空白をなくした。
"""
import argparse
import os
import shutil
import sys
import subprocess
import tempfile
from collections import namedtuple

import openpyxl
import segno
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Mm, RGBColor

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
UI_DIR = os.path.join(os.path.dirname(ROOT), "★UI研究")
COUNT_DIR = os.path.join(os.path.dirname(ROOT), "★乗降人数計算")

BASE = "https://matsumototaichi-beep.github.io/bus-ui-study/"
FONT = "Yu Gothic"
BLUE = (0x1A, 0x4B, 0x8C)
GRAY = (0x55, 0x55, 0x55)
RED = (0xB0, 0x1E, 0x1E)
HEAD_FILL = "DCE6F2"
WARN_FILL = "FBE9E7"

PAGE_W = 182          # 本文の幅 mm（A4 210 − 左右 14）
QR_MM = 46            # QR 画像の一辺 mm（余白2マス込み。読み取り部分は 4cm 以上になる）
BODY = 11.5           # 本文の文字の大きさ pt


# ---------------------------------------------------------------- アンケートURL
# UI研究の表は「回答用URL（配布用紙のQR）」、計数研究の表は「回答用URL（参加者に送る）」
URL_HEADS = ("回答用URL（配布用紙のQR）", "回答用URL（参加者に送る）")
FORMS = "https://docs.google.com/forms/"
# --draft で URL がまだのときに使う仮の URL（名前は QR で読めるように英数字にする）
DRAFT_URL = FORMS + "d/e/DRAFT-%s/viewform"
DRAFT_NAME = {"1回目_桃山台": "ui1-momoyamadai", "1回目_南口": "ui1-minamiguchi",
              "2回目_桃山台": "ui2-momoyamadai", "2回目_南口": "ui2-minamiguchi",
              "計数研究": "count"}

Survey = namedtuple("Survey", "url qsize draft")      # draft=True なら仮の URL


def survey(xlsx, name):
    """アンケートURL.xlsx から (回答用URL, 問数) を読む。URL がまだ書かれていなければ URL は None。
    URL はファイルに残さない"""
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    ws = wb["アンケートURL"]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    wb.close()
    hi = next((i for i, r in enumerate(rows) if any(h in r for h in URL_HEADS)), None)
    if hi is None or "問数" not in rows[hi]:
        sys.exit("見出しの行（%s・問数）が見つかりません: %s" % ("／".join(URL_HEADS), xlsx))
    head = rows[hi]
    ui = next(head.index(h) for h in URL_HEADS if h in head)
    qi = head.index("問数")
    for r in rows[hi + 1:]:
        first = str(r[0] or "").strip()
        if first.startswith("使わなくなったもの"):
            break                                   # ここより下（旧フォーム）は読まない
        if first != name:
            continue
        qsize = str(r[qi] or "").strip()
        if not qsize:
            sys.exit("問数が書かれていません: %s の「%s」" % (xlsx, name))
        url = str(r[ui] or "").strip()
        if not url or url.startswith("（"):        # 「（フォームを作ったら書く）」
            return None, qsize
        if not url.startswith(FORMS):
            sys.exit("アンケートの URL が読めません: %s の「%s」" % (xlsx, name))
        return url, qsize
    sys.exit("アンケートの行がありません: %s の「%s」" % (xlsx, name))


# ---------------------------------------------------------------- Word の部品
def jp(run, size, bold=False, color=None):
    run.font.size = Pt(size)
    run.bold = bold
    run.font.name = FONT
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)
    if color:
        run.font.color.rgb = RGBColor(*color)


def fmt(par, size, before=0, after=0, align=None):
    pf = par.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = Pt(round(size * 1.45, 1))     # Yu Gothic は「1行」だと行間が広すぎるので固定する
    if align is not None:
        par.alignment = align
    return par


def write(par, text, size=BODY, bold=False, color=None):
    """** で挟んだところを太字にする。\n はその場で改行する（かぎかっこの中で行が分かれないように）"""
    for li, line in enumerate(str(text).split("\n")):
        if li:
            par.add_run().add_break()
        for i, chunk in enumerate(line.split("**")):
            if chunk:
                jp(par.add_run(chunk), size, bold or i % 2 == 1, color)
    return par


def p(where, text="", size=BODY, bold=False, color=None, before=0, after=3, align=None, first=False):
    par = where.paragraphs[0] if first else where.add_paragraph()
    fmt(par, size, before, after, align)
    return write(par, text, size, bold, color)


def h2(doc, text, new_page=False, before=5, first=False):
    par = p(doc, text, size=14, bold=True, color=BLUE, before=before, after=3, first=first)
    par.paragraph_format.keep_with_next = True
    par.paragraph_format.page_break_before = new_page
    bd = OxmlElement("w:pBdr")
    bot = OxmlElement("w:bottom")
    bot.set(qn("w:val"), "single"); bot.set(qn("w:sz"), "10")
    bot.set(qn("w:space"), "2"); bot.set(qn("w:color"), "1A4B8C")
    bd.append(bot)
    par._p.get_or_add_pPr().append(bd)
    return par


def shade(cell, fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shd)


def new_table(doc, widths, grid=True):
    t = doc.add_table(rows=0, cols=len(widths))
    if grid:
        t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    t.autofit = False
    lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed")
    t._tbl.tblPr.append(lay)
    return t


def add_row(t, widths):
    cells = t.add_row().cells
    for c, w in zip(cells, widths):
        c.width = Mm(w)
    return cells


def gap(doc, size=4):
    p(doc, "", size=size, after=0)


def table(doc, body, widths, head=None, size=BODY, center=False, label_fill=HEAD_FILL):
    """1列目（head があれば見出し行も）を太字・色つきにする表。
    値がリストなら番号つきの手順、None なら左のセルとつなげる"""
    t = new_table(doc, widths)
    for ri, row in enumerate(([head] if head else []) + list(body)):
        cells = add_row(t, widths)
        is_head = head is not None and ri == 0
        for ci, val in enumerate(row):
            cells[ci].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if val is None:
                continue
            if isinstance(val, list):
                steps(cells[ci], val, size=size, first=True, after=1)
                cells[ci].paragraphs[0].paragraph_format.space_before = Pt(2)
                cells[ci].paragraphs[-1].paragraph_format.space_after = Pt(2)
                continue
            lines = str(val).split("\n")
            for li, line in enumerate(lines):
                par = p(cells[ci], line, size=size, bold=is_head or ci == 0,
                        before=2 if li == 0 else 0, after=2 if li == len(lines) - 1 else 0,
                        align=WD_ALIGN_PARAGRAPH.CENTER if center else None, first=li == 0)
                par.paragraph_format.keep_with_next = True
            if is_head or (ci == 0 and label_fill):
                shade(cells[ci], label_fill or HEAD_FILL)
        for ci in range(len(row) - 1, 0, -1):
            if row[ci] is None:
                cells[ci - 1].merge(cells[ci])
    gap(doc)
    return t


def steps(where, items, size=BODY, first=False, start=1, after=3):
    """番号つきの手順。2行目以降を番号の右にそろえる"""
    for i, text in enumerate(items, start=start):
        par = where.paragraphs[0] if (first and i == start) else where.add_paragraph()
        fmt(par, size, 0, after)
        pf = par.paragraph_format
        pf.left_indent = Mm(7)
        pf.first_line_indent = Mm(-7)
        pf.tab_stops.add_tab_stop(Mm(7))
        jp(par.add_run("%d.\t" % i), size, True, BLUE)
        write(par, text, size)


def qr_box(cell, png, title, url, draft=False):
    """右の列に QR と、その下に名前と小さく URL。仮の QR には赤で「仮」と入れる"""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    par = fmt(cell.paragraphs[0], BODY, 0, 0, WD_ALIGN_PARAGRAPH.CENTER)
    par.paragraph_format.line_spacing = None      # 画像の行は固定の行間にしない（切れるため）
    par.add_run().add_picture(png, width=Mm(QR_MM))
    for line in title.split("\n"):
        p(cell, line, size=12, bold=True, after=0, align=WD_ALIGN_PARAGRAPH.CENTER)
    if draft:
        p(cell, "仮のQR（本番では使えない）", size=10, bold=True, color=RED, after=0,
          align=WD_ALIGN_PARAGRAPH.CENTER)
    p(cell, url, size=6.5, color=GRAY, after=0, align=WD_ALIGN_PARAGRAPH.CENTER)


def with_qr(doc, png, title, url, fill, draft=False):
    """左に手順など、右に QR の2列（枠線なし）。fill(左のセル) で左を書く"""
    w = [PAGE_W - QR_MM - 8, QR_MM + 8]
    t = new_table(doc, w, grid=False)
    left, right = add_row(t, w)
    left.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    fill(left)
    qr_box(right, png, title, url, draft)
    gap(doc)


def header(doc, tag_lines):
    """左に題名、右に色つきの見分け（UI研究 1回目 など）"""
    w = [PAGE_W - 50, 50]
    t = new_table(doc, w, grid=False)
    left, right = add_row(t, w)
    left.vertical_alignment = right.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p(left, "バスの人数アプリ　実験の用紙", size=19, bold=True, after=1, first=True)
    p(left, "大和大学　情報学部　松本泰知　卒業研究", size=10.5, color=GRAY, after=0)
    shade(right, "1A4B8C")
    for i, (text, size) in enumerate(tag_lines):
        p(right, text, size=size, bold=True, color=(0xFF, 0xFF, 0xFF), before=3 if i == 0 else 0,
          after=3 if i == len(tag_lines) - 1 else 0, align=WD_ALIGN_PARAGRAPH.CENTER, first=i == 0)
    gap(doc, 2)


# ---------------------------------------------------------------- 共通の中身
def timetable(doc):
    h2(doc, "時間割　自分の時間帯に ○", before=4)
    w = [14, 20, 26, 32, 28, 34, 28]
    t = table(doc, [
        ("", "**昼**", "11:35", "**12:03**", "12:48", "**13:48**", "14:32"),
        ("", "**夕方**", "15:25", "**15:55**", "16:40", "**17:18**", "18:03"),
    ], w, head=("○", "", "集合", "ゆき\n南口発", "桃山台着", "かえり\n桃山台発", "南口着"),
        size=13, center=True, label_fill=None)
    for row in t.rows[1:]:
        shade(row.cells[1], HEAD_FILL)      # ○の列は書き込むので白のまま
    p(doc, "集合：**ＪＲ吹田駅　南口**（北口ではありません）", size=12, after=0)


# ★2026-10-08: 配布カードは1回目・計数研究では当日受け取るので、持ちものに入れるのは2回目だけ
def requests(doc, card=False):
    h2(doc, "お願い")
    table(doc, [
        ("持ちもの", "充電したスマホ　／　ICカードか小銭" + ("　／　配布カード（ID とパスワード）" if card else "")),
        ("乗るバス", "上の時刻の便だけに乗ってください"),
        ("運賃", "実費をお支払いします"),
        # 参加者向け案内の「許可していただきます」「位置情報から…把握するためです」を受けた同意の文なので、
        #   2文とも言い回しを保つ（「把握する」だけ「知る」にした。「位置情報から」は記録する目的なので残す）
        ("位置情報", "実験中は位置情報の共有を許可していただき、**記録します**。\n"
                     "位置情報からバスの便などの運行情報を知るためです。"),
        ("アプリ", "**画面の指示どおりに操作してください**"),
        ("困ったとき", "松本にチャットで連絡してください"),
        ("やめたいとき", "いつでもやめられます。松本にチャットで連絡してください"),
    ], [30, PAGE_W - 30])


# 表の「当日の流れ」で、UI研究と計数研究に共通の手順
CARD = "南口に集合。**配布カード**（ID とパスワード）を受け取る"
READ_QR = "右の**QR**をスマホのカメラで読む"
LOGIN = "「**ログイン**」タブ（「新規登録」ではない）で、カードの ID とパスワードを入れる"
# ★2026-10-08(2): iPhone の Safari では、許可を聞かれずに拒否されることがある。そのときは
#   アプリ（uiStudy-0.9.4）が押せる帯「位置情報がオフです。ここを押して直し方を見る」を出し、
#   直し方の画面もその日1回は自動で開く。同じことを用紙に書かない
GEO = "位置情報の利用を聞かれたら「**許可**」"
PRACTICE = ("のりばで練習：「**練習**」を押す → 待っている人の数を\n"
            "**2回**答える → 画面の上の「**練習モード**」をタップして終える")
BOARD = "実験者が指定したバスに乗る（のりば2、**[2]** のバス）"
SEAT = "座ったら「**着席**」、立っていたら「**立席**」を押す。変わったら押し直す"
# 「かえりに切り替える」は裏の1に書くので、表では裏へ案内するだけにする
SWITCH = "桃山台駅で降りたら、**裏**を見る"


# 桃山台で乗り間違えないために（docs/old/operation_plan.md「5. 乗り間違いを防ぐ」）
#   [2] の行先表示は「ＪＲ吹田駅（南口）」。「ＪＲ吹田駅（北口）」と出るのは [5]。
#   昼は毎時 :24 発の [3] が南口に行くが経路が違う。効くのは「指定の1本以外に乗らない」だけ
#   ★2026-10-08: いちばん見分けにくいのは、番号も行き先も同じ [2] 南口行きで時刻だけ違うバス
#   （昼は桃山台に着く 12:48 に、夕方は着いて8分後の 16:48 に出る）。1行目に分けて書く
KAERI_BUS = [
    ("乗るバス", "**13:48**（昼）／**17:18**（夕方）発\n「**[2] ＪＲ吹田駅（南口）**」行き　**だけ**"),
    ("乗らないバス", "同じ「[2] ＪＲ吹田駅（南口）」でも、**ほかの時刻**のバス\n"
                     "ほかの番号のバス（行き先が南口でも）\n「ＪＲ吹田駅（北口）」行き"),
]
KAERI_STEPS = [
    # Word では ** は太字になるだけなので、太字の前後に空白を入れない（「のりば へ」と助詞の前が空かないように）
    "かえりのバスまで、昼は**約60分**、夕方は**約38分**。\n発車までに**2番のりば**へ",
    "下のバスに乗る。**終点まで乗る**",
    # ★2026-10-08: 桃山台では実験者がいないので、かえりでも押すことを裏に書く
    #   （人数アプリはゆきの値が残る。乗降アプリは別の保存先で、未設定から始まる）
    "乗ったら「**着席**」か「**立席**」を押す",
]


def goal(doc, sv, title, answer, size=12.5, qr=None):
    """見出しも QR の左に入れる（計数研究の裏が1ページに収まるように。3種類とも同じ形にする）"""
    def fill(c):
        h2(c, "南口に着いたら", before=0, first=True)
        steps(c, [
            answer % sv.qsize,
            "答えたら解散です。おつかれさまでした",
        ], size=size, after=4)
    with_qr(doc, qr(sv.url), title, sv.url, fill, sv.draft)


# ---------------------------------------------------------------- UI研究
def ui_doc(rnd, qr, yuki, kaeri):
    """yuki＝桃山台駅で答えるアンケート、kaeri＝南口で答えるアンケート（どちらも Survey）"""
    doc = setup("配布用紙 UI研究 %d回目" % rnd)
    header(doc, [("UI研究", 15), ("%d回目" % rnd, 26)])
    timetable(doc)
    requests(doc, card=(rnd == 2))

    h2(doc, "当日の流れ")
    first = CARD if rnd == 1 else "南口に集合。1回目にもらった**配布カード**を用意する"
    app_url = BASE + "?round=%d" % rnd
    with_qr(doc, qr(app_url), "人数アプリ　%d回目" % rnd, app_url, lambda c: steps(c, [
        first, READ_QR, LOGIN,
        "「**ゆき　吹田駅 → 桃山台駅**」を選び、画面を実験者に見せる",
        GEO, PRACTICE, BOARD, SEAT,
        "バス停を発車するたびに、人数を答える（**数え方は裏**）",
        SWITCH,
    ], first=True))

    # ---- 裏（★2026-10-08(2): 桃山台駅で ゆき のアンケートに答える。手順が QR の左に入るので文字を 12.5pt にした）
    big = 12.5
    h2(doc, "桃山台駅ですること", new_page=True, before=0)
    # 「よろしいですか？」で「キャンセル」を押すと、ゆきのままになる（index.html の switchLeg）
    # 紙の QR をカメラで読むとフォームは別のタブで開き、人数アプリは裏に回る。桃山台駅には実験者がいないので、
    #   答えたあとアプリに戻ることを書く（5・6 はアプリでの操作）
    with_qr(doc, qr(yuki.url), "アンケート　ゆき", yuki.url, lambda c: steps(c, [
        "降りたら、画面右上の「**かえりに切り替える**」を押す。\n「よろしいですか？」と出たら「**OK**」",
        "かえりのバスを待つあいだに、右のQRから\n**アンケート（ゆき）**に答える（%s）。\n"
        "答えたら人数アプリに戻る" % yuki.qsize,
    ] + KAERI_STEPS + [
        "かえりのバスでも、バス停を発車するたびに人数を答える",
    ], size=big, first=True, after=4), yuki.draft)
    table(doc, KAERI_BUS, [36, PAGE_W - 36], size=big, label_fill=WARN_FILL)

    h2(doc, "数え方", before=7)
    table(doc, [
        ("数に入れる人", "車内にいる人全員。**自分も入れる**"),
        ("入れない人", "運転士"),
        ("いつ", "バス停を**発車したら**数えて答え、**次のバス停までに送る**"),
    ], [36, PAGE_W - 36], size=big)

    h2(doc, "人数の答え方", before=7)
    steps(doc, [
        "混み具合を6つから選ぶ",
        "車内の人数を入れる",
        "次に停まるバス停を選ぶ",
        "「**送信**」を押す",
    ], size=big, after=3)
    gap(doc, 6)

    goal(doc, kaeri, "アンケート　かえり", "右のQRから**アンケート（かえり）**に答える\n（%s）", size=big, qr=qr)
    return doc


# ---------------------------------------------------------------- 計数研究
def count_doc(qr, sv):
    doc = setup("配布用紙 計数研究")
    header(doc, [("計数研究", 24)])
    timetable(doc)
    requests(doc)

    h2(doc, "当日の流れ")
    yuki_url = BASE + "?study=count"
    with_qr(doc, qr(yuki_url), "ゆき　人数アプリ", yuki_url, lambda c: steps(c, [
        CARD, READ_QR, LOGIN,
        "画面の上に「**計数研究　ゆき**」と出たら、実験者に見せる",
        GEO, PRACTICE, BOARD, SEAT,
        "バス停を発車するたびに、車内の人数を**1から数えて**答える（**数え方は裏**）",
        SWITCH,
    ], first=True))

    # ---- 裏
    h2(doc, "桃山台駅ですること", new_page=True, before=0)
    kaeri_url = BASE + "counter.html"
    # ★2026-10-08: 左の列は狭いので、ボタン名やかぎかっこが行の途中で分かれない長さにする。
    #   乗降アプリは乗った人・降りた人の足し引きなので、最初の桃山台駅で数え損ねると終点までずれる。
    #   桃山台駅で乗ってきた人（自分たちも）を数えることを手順に入れる
    with_qr(doc, qr(kaeri_url), "かえり　乗降アプリ", kaeri_url, lambda c: steps(c, [
        "降りたら、画面右上の「**かえりに切り替える**」を押す",
        "「乗降アプリを開きますか？」と出たら「**OK**」。\n「**計数研究　かえり**」が開く（開かなければ右のQRで）",
        "ログイン画面が出たら、カードの ID とパスワードを入れる",
    ] + KAERI_STEPS + [
        "**桃山台駅で乗ってきた人**（自分たちも）を数えて「**送信**」",
    ], size=12, first=True))
    table(doc, KAERI_BUS, [34, PAGE_W - 34], size=12, label_fill=WARN_FILL)

    # ゆき（人数アプリ）と かえり（乗降アプリ）を横に並べて、違いが一目で分かるようにする
    h2(doc, "数え方・答え方")
    table(doc, [
        ("数える人", "車内にいる人全員。**自分も入れる**",
         "**乗ってきた人・降りた人**。自分たちも数える（桃山台駅で乗るとき・終点で降りるとき）"),
        ("数えない人", "運転士", None),
        ("いつ", "**発車したら**、**1から数えて**答え、\n次のバス停までに送る",
         "バス停に**停車したら**数える"),
        ("答え方", ["混み具合を6つから選ぶ", "車内の人数を入れる", "次に停まるバス停を選ぶ",
                    "「**送信**」を押す"],
         ["バス停を確かめる。違ったら「**変える**」",
          "「**乗ってきた人**」1人ごとに「**＋1**」",
          "「**降りた人**」も同じ。押しすぎたら「**−1**」",
          "「**送信**」→ 次のバス停へ",
          "終点でも降りた人を数えて送る\n「**終点です。おつかれさまでした**」と出る"]),
    ], [25, 65, PAGE_W - 90], head=("", "ゆき（人数アプリ）", "かえり（乗降アプリ）"), size=11.5)

    gap(doc, 4)
    goal(doc, sv, "アンケート　計数研究", "右のQRから**アンケート**に答える（%s）", size=12.5, qr=qr)
    return doc


def setup(title):
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = FONT
    st.font.size = Pt(BODY)
    st.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)
    st.paragraph_format.space_after = Pt(0)
    for s in doc.sections:
        s.page_width, s.page_height = Mm(210), Mm(297)
        s.top_margin = s.bottom_margin = Mm(11)
        s.left_margin = s.right_margin = Mm(14)
        s.header_distance = s.footer_distance = Mm(6)
    doc.core_properties.title = title
    doc.core_properties.author = "松本泰知"
    return doc


# ---------------------------------------------------------------- PDF（Word で書き出す）
PS = r"""
$ErrorActionPreference = 'Stop'
$list = Get-Content -Encoding UTF8 -LiteralPath $env:HANDOUT_LIST | Where-Object { $_ -ne '' }
$word = New-Object -ComObject Word.Application
$word.Visible = $false
$word.DisplayAlerts = 0
try {
  $i = 0
  foreach ($path in $list) {
    $doc = $word.Documents.Open($path, $false, $true)
    $doc.Repaginate()
    $pages = $doc.ComputeStatistics(2)
    $pdf = [System.IO.Path]::ChangeExtension($path, '.pdf')
    $doc.ExportAsFixedFormat($pdf, 17)
    $doc.Close($false)
    Write-Output ("PAGES " + $i + " " + $pages)
    $i++
  }
} finally {
  $word.Quit()
  [void][System.Runtime.InteropServices.Marshal]::ReleaseComObject($word)
}
"""


def export_pdf(paths):
    """Word で開いてページ数を数え、PDF に書き出す。ページ数のリストを返す"""
    with tempfile.TemporaryDirectory() as tmp:
        lst = os.path.join(tmp, "list.txt")
        with open(lst, "w", encoding="utf-8") as f:
            f.write("\n".join(paths) + "\n")
        env = dict(os.environ, HANDOUT_LIST=lst)
        r = subprocess.run(["powershell", "-NoProfile", "-NonInteractive", "-Command", PS],
                           capture_output=True, env=env)
    out = r.stdout.decode("cp932", "replace")
    if r.returncode != 0:
        sys.exit("Word で PDF にできませんでした（書き出し先のファイルは変えていません）\n" + out
                 + r.stderr.decode("cp932", "replace"))
    pages = {}
    for line in out.splitlines():
        if line.startswith("PAGES "):
            _, i, n = line.split()
            pages[int(i)] = int(n)
    return [pages.get(i) for i in range(len(paths))]


# ---------------------------------------------------------------- 実行
def inside(path, folder):
    """path が folder そのもの、または folder の中なら True（大文字・小文字と区切りの違いは見ない）"""
    path, folder = (os.path.normcase(os.path.realpath(x)) for x in (path, folder))
    try:
        return os.path.commonpath([path, folder]) == folder
    except ValueError:              # ドライブが違う
        return False


def main():
    ap = argparse.ArgumentParser(description="当日に配る用紙（A4 両面1枚）を Word と PDF にする")
    ap.add_argument("--draft", metavar="フォルダ",
                    help="下書き。URL がないアンケートは仮の QR にして、3枚ともこのフォルダに書き出す")
    a = ap.parse_args()

    out_ui, out_ct = UI_DIR, COUNT_DIR
    if a.draft:
        out_ui = out_ct = os.path.abspath(a.draft)
        real = {os.path.normcase(os.path.abspath(d)) for d in (UI_DIR, COUNT_DIR)}
        if os.path.normcase(out_ui) in real:
            sys.exit("--draft には本番のフォルダ（★UI研究・★乗降人数計算）以外を指定してください")
        # ★2026-10-08(3): 下書きにも本物のアンケート URL が入る（URL がそろったもの）。public のリポジトリの中には置かない
        if inside(out_ui, ROOT):
            sys.exit("--draft にはリポジトリ（%s）の外のフォルダを指定してください。"
                     "下書きにも本物のアンケートの URL が入るためです" % os.path.basename(ROOT))
        os.makedirs(out_ui, exist_ok=True)

    def get(xlsx, name):
        url, qsize = survey(xlsx, name)
        if url:
            return Survey(url, qsize, False)
        if a.draft:
            return Survey(DRAFT_URL % DRAFT_NAME[name], qsize, True)
        return None

    ui_x = os.path.join(UI_DIR, "アンケートURL.xlsx")
    ct_x = os.path.join(COUNT_DIR, "アンケートURL.xlsx")
    plans = []          # (書き出す docx, 作る関数, 仮の QR があるか)
    for rnd in (1, 2):
        names = ["%d回目_桃山台" % rnd, "%d回目_南口" % rnd]
        sv = [get(ui_x, n) for n in names]
        path = os.path.join(out_ui, "配布用紙_UI研究_%d回目.docx" % rnd)
        if None in sv:
            print("書き出しません: %s（アンケートの URL がまだ: %s）"
                  % (os.path.basename(path)[:-5], "、".join(n for n, s in zip(names, sv) if s is None)))
            continue
        plans.append((path, lambda qr, rnd=rnd, sv=sv: ui_doc(rnd, qr, *sv), any(s.draft for s in sv)))
    sc = get(ct_x, "計数研究")
    path = os.path.join(out_ct, "配布用紙_計数研究.docx")
    if sc is None:
        print("書き出しません: 配布用紙_計数研究（アンケートの URL がまだ: 計数研究）")
    else:
        plans.append((path, lambda qr: count_doc(qr, sc), sc.draft))
    if not plans:
        sys.exit("書き出せる用紙がありません。アンケートURL.xlsx に URL を書くか、--draft で下書きにしてください")

    # ★2026-10-08(3): docx と PDF は一時フォルダで作り、3枚とも2ページだと確かめてから書き出し先へ写す。
    #   2ページに収まらないときは、書き出し先にある前の版（印刷に使える版）に触らずに止まる
    with tempfile.TemporaryDirectory() as tmp:
        made = {}

        def qr(url):
            if url not in made:
                path = os.path.join(tmp, "qr%d.png" % len(made))
                segno.make(url, error="m").save(path, scale=20, border=2)
                made[url] = path
            return made[url]

        work = []
        for path, make, _ in plans:
            w = os.path.join(tmp, os.path.basename(path))
            make(qr).save(w)
            work.append(w)

        pages = export_pdf(work)
        bad = False
        for (path, _, draft), n in zip(plans, pages):
            print("%s  %s ページ%s" % (os.path.basename(path)[:-5] + "（.docx / .pdf）", n,
                                     "　仮のQRあり" if draft else ""))
            bad = bad or n != 2
        if bad:
            sys.exit("2ページ（表・裏）に収まっていません。文字を減らすか小さくしてください"
                     "（書き出し先のファイルは変えていません）")

        pairs = [(w[:-5] + ext, path[:-5] + ext)
                 for (path, _, _), w in zip(plans, work) for ext in (".docx", ".pdf")]
        for _, dst in pairs:                 # 開いたままのファイルがあれば、1つも写さずに止まる
            if os.path.exists(dst):
                try:
                    open(dst, "r+b").close()
                except OSError:
                    sys.exit("%s を開いたままなら閉じてください（書き出し先のファイルは変えていません）" % dst)
        for src, dst in pairs:
            shutil.copyfile(src, dst)

    paths = [pl[0] for pl in plans]
    print("書き出しました: " + "、".join(sorted({os.path.dirname(x) for x in paths})))


if __name__ == "__main__":
    main()
