# -*- coding: utf-8 -*-
"""UI研究のアンケート（Google フォーム4つ）を作る

    python tools/make_forms.py --bank --force   設問一覧.xlsx を作り直す（このファイルの ROWS に戻る）
    python tools/make_forms.py                  Excel から GoogleForms生成.gs を作る

設問の正本は docs/questionnaire/設問一覧.xlsx。文言を直すときはそこを直す。

1回の往復で2回答える。桃山台駅で降りたらゆきの分、南口に着いたらかえりの分。
  1回目_桃山台（ゆき・数字キー）   … 共通5問 ＋ 自由記述（6問）
  1回目_南口（かえり・＋−ボタン）  … 共通5問 ＋ SUS 10問 ＋ 自由記述（16問）
  2回目_桃山台（ゆき・＋−ボタン）  … 共通5問 ＋ 自由記述（6問）
  2回目_南口（かえり・数字キー）   … 共通5問 ＋ SUS 10問 ＋「わからない」ボタン3問 ＋ 自由記述（19問）
SUS は System Usability Scale（使いやすさを測る決まった10問）。
回答は1つのスプレッドシートの、フォームと同じ名前のシート4枚にたまる（.gs の createAll が作る）。
"""
import os, sys
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUTDIR = os.path.join(ROOT, "docs", "questionnaire")
XLSX = os.path.join(OUTDIR, "設問一覧.xlsx")
GS = os.path.join(OUTDIR, "GoogleForms生成.gs")

# ★2026-10-07:「何が分かるか」列を足した。研究で知りたいこと（入れ方・わからないボタン・混雑）の
#   どれに使う設問かを1行で書く。フォームと紙には出ない。
COLS = ["用紙", "節", "番号", "形式", "設問文", "選択肢（／区切り）", "必須", "補足（説明文）", "何が分かるか"]
# ★2026-10-08:「回答後」を足した。送信したあとに出る文（設問文の列に書く）。
KATA = ["表紙", "説明", "ページ", "ID", "短文", "段落", "単一選択", "複数選択", "5段階", "日付", "回答後"]

# 用紙の名前は、フォーム・回答のシート・アンケートURL.xlsx の行の名前と一字一句同じ
R1M, R1S, R2M, R2S = "1回目_桃山台", "1回目_南口", "2回目_桃山台", "2回目_南口"
YOUSHI = [R1M, R1S, R2M, R2S]
# .gs の関数名（_ で終わる関数は Apps Script の実行メニューに出ない。出るのは whoami と createAll だけ）
FN = {R1M: "make1Momoyamadai_", R1S: "make1Minamiguchi_",
      R2M: "make2Momoyamadai_", R2S: "make2Minamiguchi_"}

YUKI, KAERI = "ゆき", "かえり"
# 入れ方の名前と、どの画面かの説明。条件を明かさないよう、その入れ方を使い終えた区間のあとでだけ出す
SUJI = ("数字キー", "0〜9 のキーで人数を打つ画面")
PM = ("＋−ボタン", "−5・−1・＋1・＋5 で人数を合わせる画面")

# 5段階はすべて「1＝少ない・弱い、5＝多い・強い」の向きにそろえる
# ★2026-10-07:「混んでいましたか」「答えやすかったですか」と片側だけで聞くと「はい」側に寄るので、
#   問2・問5 は「〜の混み具合」「〜の答えやすさ」と聞き、向きは両端のラベルだけで示す。
YASUI = "答えにくかった／答えやすかった"
KONDA = "空いていた／混んでいた"
HINDO = "なかった／何度もあった"
# ★2026-10-07: 選択肢は「ボタンがあった回」ではなく回の番号にする。ボタンが答えの理由だと匂わせず、
#   先頭に出る選択肢がボタンのある回に偏らない（組で逆になる）。覚え違いがあっても組から実際の条件に戻せる。
DOTCHI = "1回目／2回目／変わらない"
LIKERT = "まったくそう思わない／とてもそう思う"

# ★2026-10-08: SUS は1回目・2回目の両方の南口で聞く。「わからない」ボタンがある回とない回の使いやすさを、
#   同じ物差しで比べるため。決まった尺度なので、10問の意味・順番・5段階の向きは変えない。
#   点数の出し方が決まっているので、10問すべて必須にする。
SUS = [
    "このアプリを何度も使いたいと思う",
    "このアプリは必要以上に複雑だと思う",
    "このアプリは簡単に使えると思う",
    "このアプリを使うには、詳しい人の助けが必要だと思う",
    "このアプリのいろいろな機能は、うまくまとまっていると思う",
    "このアプリには一貫性のないところが多いと思う",
    "たいていの人は、このアプリの使い方をすぐ覚えられると思う",
    "このアプリはとても使いにくいと思う",
    "このアプリを自信をもって使えた",
    "このアプリを使う前に、いろいろ覚える必要があった",
]

# ★2026-10-07: 88回答えた直後に長いアンケートは重すぎる（監査）。研究で知りたい3つ
#   （入れ方・わからないボタン・混雑）に対応しない設問と、アプリの記録で分かる設問は入れない。
#   任意の設問は面倒なものだけ抜けて集計が偏るので、すべて必須にする。
# ★2026-10-08: 区間ごとに答える形にした。往復のあとにまとめて聞くと、ゆきとかえりの記憶が混ざる。
#   問1〜5 は4つのフォームで同じにし、区間（ゆき／かえり）と入れ方だけを差し込む。
#   入れ方は、その区間で使ったものだけを聞く。1回目はゆき＝数字キー・かえり＝＋−ボタン、2回目はその逆なので、
#   1人が両方の入れ方を、ゆきとかえりで1回ずつ評価する。区間や順番（あとで使ったほうが慣れていた）の影響と分けられる。
#   「わからない」ボタンは組によって1回目にあったり2回目にあったりするので、
#   両方の回を終えたあとの 2回目_南口 でだけ、SUS のあとに聞く。それより前のフォームには一切出さない。


def kyotsu(youshi, kukan, irekata):
    """問1〜5。4つのフォームで同じ。区間と入れ方だけが変わる"""
    name, gamen = irekata
    return [
        (youshi, "はじめに", "問1", "ID", "カードに書いてあるID（例：A）", "", "○", "",
         "アンケートとアプリの記録をつなぐ"),
        # ★2026-10-08: 区間の名前を問2〜5 に入れる。往復全体や前の回のことを思い出して答えないように。
        #   ★2026-10-09（松本さんの決定）: 問4 にも入れた。南口でゆきの分まで含めて答えるのを防ぐ
        (youshi, "混み具合", "問2", "5段階", "%sのバスの混み具合" % kukan, KONDA, "○", "",
         "その区間の混み具合。同じ区間のクラスの回答・問5 と並べて見る"),
        # ★2026-10-07:「混んでいて、」は付けない。立っていた・降りる人が多かったなど別の理由で数えきれなかった人が
        #   答えられず、混雑との関係も答える側に決めさせてしまうため。混雑との関係は問2 と並べて見る。
        (youshi, "混み具合", "問3", "5段階", "%sのバスで人数を数えきれないことがありましたか" % kukan, HINDO, "○", "",
         "数えきれないことがどれだけあったか。問2 と並べて、混んでいると数えにくいかを見る"),
        (youshi, "混み具合", "問4", "5段階", "%sのバスで、数えきれないまま、だいたいの人数を入れて送ったことがありましたか" % kukan, HINDO, "○", "",
         "数えられないときに人数を入れてしまったか（記録では分からない）。ボタンがあった回となかった回で比べる"),
        (youshi, "入れ方", "問5", "5段階", "%sで使った%sの答えやすさ" % (kukan, name), YASUI, "○", gamen,
         "その区間で使った入れ方の答えやすさ。同じ人の数字キーと＋−ボタンを比べる"),
    ]


def jiyu(youshi, no):
    return [(youshi, "最後に", no, "段落", "答えにくかったことがあれば書いてください（なければ「なし」）", "", "○", "",
             "点数の理由")]


def hyoshi(youshi, kai, kukan, doko, fun):
    return [(youshi, "表紙", "", "表紙", "バスの人数アプリ　%s　%sのアンケート" % (kai, kukan), "", "",
             "%s答えてください（%s）。思ったとおりに選んでください。" % (doko, fun), "")]


def kaitougo(youshi, bun):
    return [(youshi, "回答後", "", "回答後", bun, "", "", "", "")]


def sus(youshi, start):
    return (
        [(youshi, "使いやすさ", "", "ページ", "アプリ全体について", "", "",
          "今回の往復で使ってみた、アプリ全体の印象をお答えください。", "")]
        + [(youshi, "使いやすさ", "問%d" % (start + i), "5段階", t, LIKERT, "○", "",
            "アプリ全体の使いやすさの点数（SUS＝使いやすさを測る決まった10問）。ボタンがあった回となかった回で比べる"
            if i == 0 else "同上")
           for i, t in enumerate(SUS)]
    )


MOMO = "桃山台駅で降りたら"
MINAMI = "南口に着いたら"
# 桃山台のあとはかえりがあるので続けてもらう。かえりの入れ方（条件）は書かない
OWARI_MOMO = "ありがとうございました。かえりのバスでも、バス停を発車するたびに人数を答えてください。"

ROWS = (
    # ============================== 1回目_桃山台 ==============================
    hyoshi(R1M, "1回目", YUKI, MOMO, "1分ほど")
    + kyotsu(R1M, YUKI, SUJI)
    + jiyu(R1M, "問6")
    + kaitougo(R1M, OWARI_MOMO)

    # ============================== 1回目_南口 ==============================
    + hyoshi(R1S, "1回目", KAERI, MINAMI, "4分ほど")
    + kyotsu(R1S, KAERI, PM)
    + sus(R1S, 6)
    + jiyu(R1S, "問16")
    # 1回目の人は別の日に2回目があるので「今回は」
    + kaitougo(R1S, "ご協力ありがとうございました。今回はこれで終わりです。")

    # ============================== 2回目_桃山台 ==============================
    + hyoshi(R2M, "2回目", YUKI, MOMO, "1分ほど")
    + kyotsu(R2M, YUKI, PM)
    + jiyu(R2M, "問6")
    + kaitougo(R2M, OWARI_MOMO)

    # ============================== 2回目_南口 ==============================
    + hyoshi(R2S, "2回目", KAERI, MINAMI, "5分ほど")
    + kyotsu(R2S, KAERI, SUJI)
    # ★2026-10-09（松本さんの決定）: SUS を先にする。1回目_南口 と同じく問5 のすぐあと（問6〜15）に置き、
    #   ボタンの話を聞く前に答えてもらう（SUS の点数がボタンの設問に引っぱられないように）。
    #   問16〜18 が「わからない」ボタンの3問、問19 が自由記述。
    + sus(R2S, 6)
    + [
        # ★2026-10-07: 共通の問と SUS を答え終えてから見せるよう、ページを分ける。
        #   どちらの回にボタンがあったかは書かない（組によって逆なので決めつけない）。
        (R2S, "わからない", "", "ページ", "「わからない」ボタンについて", "", "",
         "1回目と2回目のどちらか一方だけ、人数を入れるところの下に「わからない」ボタンがありました。", ""),
        (R2S, "わからない", "問16", "単一選択", "「わからない」ボタンがあったのは、どちらの回ですか",
         "1回目／2回目／覚えていない", "○", "",
         "ボタンに気づいていたか"),
        (R2S, "わからない", "問17", "単一選択", "人数を答えやすかったのは、どちらの回ですか", DOTCHI, "○", "",
         "ボタンがあると答えやすいか（組から、ボタンがあった回を選んだかを見る）"),
        # ★2026-10-07:「正直に答えられた」は、もう一方の回は正直でなかったと認めさせる聞き方で「変わらない」に寄る。
        #   ボタンの狙い（数えられないときに数を作らない）をそのまま言葉にした。
        (R2S, "わからない", "問18", "単一選択", "数えたとおりの人数を答えられたのは、どちらの回ですか", DOTCHI, "○", "",
         "ボタンがあると正直に答えられるか（同上）"),
    ]
    + jiyu(R2S, "問19")
    + kaitougo(R2S, "ご協力ありがとうございました。これで終わりです。")
)


# ---------------------------------------------------------------- Excel を作る
def write_bank(force=False):
    if os.path.exists(XLSX) and not force:
        sys.exit("既に %s があります。上書きするなら --force を付けてください。" % XLSX)
    os.makedirs(OUTDIR, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "設問"
    ws["A1"] = "UI研究のアンケート　設問一覧"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = ("直したら python tools/make_forms.py（フォームを作るスクリプト）と "
                "python tools/make_questionnaire.py（紙の予備）を流す。")
    ws["A2"].font = Font(size=9, color="555555")

    thin = Side(style="thin", color="AAAAAA")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    for i, c in enumerate(COLS, start=1):
        cell = ws.cell(row=4, column=i, value=c)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DCE6F1")
        cell.border = box
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for i, w in enumerate([14, 11, 7, 10, 48, 36, 6, 38, 44], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    for r, row in enumerate(ROWS, start=5):
        for c, v in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.border = box
            cell.alignment = Alignment(vertical="top", wrap_text=(c in (5, 6, 8, 9)))
        if row[3] in ("表紙", "説明", "ページ", "回答後"):
            for c in range(1, len(COLS) + 1):
                ws.cell(row=r, column=c).fill = PatternFill("solid", fgColor="F2F2F2")

    dv = DataValidation(type="list", formula1='"%s"' % ",".join(KATA), allow_blank=True)
    ws.add_data_validation(dv)
    dv.add("D5:D%d" % (4 + len(ROWS)))
    ws.freeze_panes = "A5"

    wb.save(XLSX)
    print("書き出しました: " + XLSX)


def read_bank():
    """設問一覧.xlsx の行（9列の文字列）。make_questionnaire.py も使う"""
    if not os.path.exists(XLSX):
        sys.exit("%s がありません。先に python tools/make_forms.py --bank を流してください。" % XLSX)
    ws = load_workbook(XLSX)["設問"]
    rows = []
    for r in range(5, ws.max_row + 1):
        vals = [ws.cell(row=r, column=c).value for c in range(1, len(COLS) + 1)]
        if not any(vals):
            continue
        row = [("" if v is None else str(v)) for v in vals]
        if row[0] not in YOUSHI:
            sys.exit("設問一覧.xlsx の %d 行目：用紙「%s」は %s のどれでもありません。" % (r, row[0], "／".join(YOUSHI)))
        rows.append(row)
    for y in YOUSHI:
        if not [x for x in rows if x[0] == y and x[3] == "表紙"]:
            sys.exit("%s の「表紙」の行がありません。" % y)
    return rows


def count_q(rows, youshi):
    return len([r for r in rows if r[0] == youshi and r[2]])


# ---------------------------------------------------------------- .gs を作る
def js(s):
    return "'" + str(s).replace("\\", "\\\\").replace("'", "\\'").replace("\n", "\\n") + "'"


def emit_item(row, out):
    youshi, setsu, no, kata, bun, sel, must, hosoku = row[:8]
    label = ("%s %s" % (no, bun)).strip()
    choices = [x.strip() for x in (sel or "").split("／") if x.strip()]

    if kata in ("表紙", "回答後"):
        return                      # フォームの題名・説明文・回答後の文に使う（form_fn）
    elif kata == "ページ":
        out.append("  it = f.addPageBreakItem().setTitle(%s);" % js(bun))
        if hosoku:
            out.append("  it.setHelpText(%s);" % js(hosoku))
    elif kata == "説明":
        out.append("  it = f.addSectionHeaderItem().setTitle(%s);" % js(bun))
        if hosoku:
            out.append("  it.setHelpText(%s);" % js(hosoku))
    elif kata == "ID":
        # ★2026-10-07: IDを打ち間違えるとアプリの記録とつなげないので、A〜X の1文字だけ受け付ける
        out.append("  it = f.addTextItem().setTitle(%s);" % js(label))
        out.append("  it.setValidation(FormApp.createTextValidation().setHelpText(%s)"
                   ".requireTextMatchesPattern(%s).build());"
                   % (js("カードのIDを1文字で入れてください（例：A）"), js(r"^\s*[A-Xa-xＡ-Ｘａ-ｘ]\s*$")))
    elif kata == "短文":
        out.append("  it = f.addTextItem().setTitle(%s);" % js(label))
    elif kata == "段落":
        out.append("  it = f.addParagraphTextItem().setTitle(%s);" % js(label))
    elif kata == "日付":
        out.append("  it = f.addDateItem().setTitle(%s);" % js(label))
    elif kata == "単一選択":
        out.append("  it = f.addMultipleChoiceItem().setTitle(%s).setChoiceValues([%s]);"
                   % (js(label), ", ".join(js(c) for c in choices)))
    elif kata == "複数選択":
        out.append("  it = f.addCheckboxItem().setTitle(%s).setChoiceValues([%s]);"
                   % (js(label), ", ".join(js(c) for c in choices)))
    elif kata == "5段階":
        left = choices[0] if choices else "1"
        right = choices[1] if len(choices) > 1 else "5"
        out.append("  it = f.addScaleItem().setTitle(%s).setBounds(1, 5).setLabels(%s, %s);"
                   % (js(label), js(left), js(right)))
    else:
        return
    if kata not in ("ページ", "説明"):
        if hosoku:
            out.append("  it.setHelpText(%s);" % js(hosoku))
        out.append("  it.setRequired(%s);" % ("true" if str(must).strip() == "○" else "false"))


HEAD = [
    "// UI研究のアンケート（Google フォーム4つ）と、回答がたまるスプレッドシートを作るスクリプト",
    "// docs/questionnaire/設問一覧.xlsx から tools/make_forms.py で作る。設問は Excel を直して作り直す。",
    "//",
    "// 1. Google ドライブ →「新規 → その他 → Google Apps Script」",
    "// 2. この中身をぜんぶ貼り付けて保存",
    "// 3. 関数 whoami を実行し、ログに出たアカウントを確かめる（フォームとシートはこのアカウントのドライブにできる）",
    "// 4. 関数 createAll を実行（1〜2分かかる）",
    "// 5. 実行ログの「==== ここから下を…」の行から下をすべてコピーしてチャットに貼る",
    "// 6. 回答用URLを Google にログインしていないブラウザで開き、答えられるか確かめる",
    "//",
    "// ★createAll を2回実行すると、フォーム4つと回答シートがもう1組できる（いらないほうはゴミ箱に入れる）。",
    "// ★前に作った2つのフォーム（1回目・2回目のアンケート）と回答表「事後アンケート回答」は使わない。",
    "",
    "function whoami() {",
    "  Logger.log('いま動いているアカウント: ' + Session.getEffectiveUser().getEmail());",
    "}",
    "",
]

# createAll と、回答シートを探す・並べる関数。名前（NAMES）は Python から差し込む
CREATE_ALL = r"""
var NAMES = %(names)s;
var SHEET_TITLE = %(sheet_title)s;

function createAll() {
  var makers = [%(makers)s];
  var forms = [], ss = null;
  try {
    // 作る関数は FormApp.create の直後にフォームを forms に入れる。途中で止まっても、作りかけを下のログに出せる
    for (var i = 0; i < makers.length; i++) {
      makers[i](forms);
    }
    ss = SpreadsheetApp.create(SHEET_TITLE);
    for (var j = 0; j < forms.length; j++) {
      forms[j].setDestination(FormApp.DestinationType.SPREADSHEET, ss.getId());
      var sh = findSheet_(ss.getId(), forms[j]);
      if (!sh) {
        throw new Error('「' + NAMES[j] + '」の回答シートが見つかりません');
      }
      sh.setName(NAMES[j]);
    }
    tidySheets_(ss.getId());
  } catch (e) {
    var msg = ['途中で止まりました: ' + e.message,
               '作りかけのものはゴミ箱に入れてから、もう一度 createAll を実行してください。'];
    for (var k = 0; k < forms.length; k++) {
      msg.push('  ' + NAMES[k] + '（作りかけ） ' + forms[k].getEditUrl());
    }
    if (ss) {
      msg.push('  回答シート（作りかけ） ' + ss.getUrl());
    }
    Logger.log(msg.join('\n'));
    throw e;
  }
  // 1回の Logger.log にまとめる（URL の行ごとにログの時刻が付かないように）。区切りはタブ
  var lines = ['==== ここから下をすべてコピーしてチャットに貼ってください ===='];
  for (var n = 0; n < forms.length; n++) {
    lines.push(NAMES[n] + '\t' + forms[n].getPublishedUrl() + '\t' + forms[n].getEditUrl());
  }
  lines.push('回答シート\t' + ss.getUrl());
  Logger.log(lines.join('\n'));
}

// 回答先にしたときに自動でできるシートを探す。すぐには見えないことがあるので、数回待つ
function findSheet_(ssId, form) {
  for (var n = 0; n < 10; n++) {
    SpreadsheetApp.flush();
    var sheets = SpreadsheetApp.openById(ssId).getSheets();
    for (var i = 0; i < sheets.length; i++) {
      if (isSheetOf_(sheets[i], form)) {
        return sheets[i];
      }
    }
    Utilities.sleep(2000);
  }
  return null;
}

// シートにつながったフォームの URL に、フォームの ID（または回答用URLの ID）が含まれるか
function isSheetOf_(sheet, form) {
  var url = sheet.getFormUrl();
  if (!url) {
    return false;
  }
  if (url.indexOf(form.getId()) >= 0) {
    return true;
  }
  var m = form.getPublishedUrl().match(/\/forms\/d\/e\/([^\/?#]+)/);
  return !!(m && url.indexOf(m[1]) >= 0);
}

// シートを NAMES の順に並べ、最初からある空のシート（「シート1」など）を消す。
// 4枚がそろってから消すので、最後の1枚を消すことにはならない
function tidySheets_(ssId) {
  SpreadsheetApp.flush();
  var ss = SpreadsheetApp.openById(ssId);
  for (var i = 0; i < NAMES.length; i++) {
    ss.setActiveSheet(ss.getSheetByName(NAMES[i]));
    ss.moveActiveSheet(i + 1);
  }
  var sheets = ss.getSheets();
  for (var j = 0; j < sheets.length; j++) {
    var sh = sheets[j];
    if (NAMES.indexOf(sh.getName()) < 0 && !sh.getFormUrl() && ss.getSheets().length > 1) {
      ss.deleteSheet(sh);
    }
  }
}
"""


def write_gs():
    rows = read_bank()
    out = list(HEAD)
    out += (CREATE_ALL % {
        "names": "[" + ", ".join(js(y) for y in YOUSHI) + "]",
        "sheet_title": js("UI研究 アンケート回答"),
        "makers": ", ".join(FN[y] for y in YOUSHI),
    }).strip("\n").split("\n")
    out.append("")

    for youshi in YOUSHI:
        hy = [r for r in rows if r[0] == youshi and r[3] == "表紙"][0]
        ok = [r for r in rows if r[0] == youshi and r[3] == "回答後"]
        o = ["// " + youshi,
             "function %s(forms) {" % FN[youshi],
             "  var f = FormApp.create(%s);" % js(hy[4]),
             "  forms.push(f);",
             "  f.setDescription(%s);" % js(hy[7])]
        if ok and ok[0][4]:
            o.append("  f.setConfirmationMessage(%s);" % js(ok[0][4]))
        o += ["  try { f.setCollectEmail(false); } catch (e) {}",
              # ★2026-10-07: 大学アカウントで作ると「大学の人だけ回答できる」になることがある。
              #   参加者はスマホで大学アカウントにログインしているとは限らないので外す。
              #   公開状態の設定が無い・効かない環境もあるので try で囲み、手順6で実際に開いて確かめる。
              "  try { f.setRequireLogin(false); } catch (e) {}",
              "  try { f.setPublished(true); } catch (e) {}",
              "  var it;"]
        for row in rows:
            if row[0] == youshi:
                emit_item(row, o)
        o += ["}", ""]
        out += o

    with open(GS, "w", encoding="utf-8", newline="\n") as f:
        f.write("\n".join(out))
    print("書き出しました: " + GS)
    print("  " + " ／ ".join("%s %d問" % (y, count_q(rows, y)) for y in YOUSHI))


if __name__ == "__main__":
    if "--bank" in sys.argv:
        write_bank(force=("--force" in sys.argv))
    else:
        write_gs()
