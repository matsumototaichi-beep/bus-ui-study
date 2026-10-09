# -*- coding: utf-8 -*-
"""実験の進め方を Word にする

    python tools/make_operation_docx.py

docs/実験の進め方.docx が出る。** で挟んだところが太字になる。
"""
import os
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Mm, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "docs", "実験の進め方.docx")
FONT = "Yu Gothic"


def jp(run, size, bold=False, color=None):
    run.font.size = Pt(size)
    run.bold = bold
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    if color:
        run.font.color.rgb = RGBColor(*color)


def p(doc, text="", size=10, bold=False, before=0, after=4):
    par = doc.add_paragraph()
    par.paragraph_format.space_before = Pt(before)
    par.paragraph_format.space_after = Pt(after)
    for i, chunk in enumerate(text.split("**")):
        if chunk:
            jp(par.add_run(chunk), size, bold or i % 2 == 1)
    return par


def h1(doc, text):
    p(doc, text, size=17, bold=True, after=10)


def h2(doc, text):
    par = p(doc, text, size=13, bold=True, before=14, after=4)
    par.paragraph_format.keep_with_next = True
    bd = OxmlElement("w:pBdr")
    bot = OxmlElement("w:bottom")
    bot.set(qn("w:val"), "single"); bot.set(qn("w:sz"), "8")
    bot.set(qn("w:space"), "2"); bot.set(qn("w:color"), "1A4B8C")
    bd.append(bot)
    par._p.get_or_add_pPr().append(bd)


def h3(doc, text):
    p(doc, text, size=11, bold=True, before=8, after=3).paragraph_format.keep_with_next = True


def table(doc, head, body, widths=None):
    t = doc.add_table(rows=0, cols=len(head or body[0]))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    t.autofit = False
    rows = ([head] if head else []) + body
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for ci, val in enumerate(row):
            if widths:
                cells[ci].width = Mm(widths[ci])
            par = cells[ci].paragraphs[0]
            par.paragraph_format.space_before = Pt(2)
            par.paragraph_format.space_after = Pt(2)
            par.paragraph_format.keep_with_next = ri < len(rows) - 1   # ★2026-10-07: 表がページをまたがないように
            is_head = head and ri == 0
            for i, chunk in enumerate(str(val).split("**")):
                if chunk:
                    jp(par.add_run(chunk), 9.5, is_head or i % 2 == 1)
    p(doc, "", size=4, after=0)
    return t


def steps(doc, items):
    for i, t in enumerate(items, start=1):
        par = doc.add_paragraph()
        par.paragraph_format.space_before = Pt(0)
        par.paragraph_format.space_after = Pt(3)
        par.paragraph_format.left_indent = Mm(5)
        jp(par.add_run("%d.  " % i), 10, True, (0x1A, 0x4B, 0x8C))
        for j, chunk in enumerate(t.split("**")):
            if chunk:
                jp(par.add_run(chunk), 10, j % 2 == 1)


doc = Document()
st = doc.styles["Normal"]
st.font.name = FONT
st.font.size = Pt(10)
st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
for s in doc.sections:
    s.page_width, s.page_height = Mm(210), Mm(297)
    s.top_margin = s.bottom_margin = Mm(17)
    s.left_margin = s.right_margin = Mm(17)

h1(doc, "実験の進め方")

# ---------------------------------------------------------------- 全体
h2(doc, "研究は2つある")
# ★2026-10-08: UI研究のアンケートは区間ごとに2つ（桃山台駅でゆき・南口でかえり）。HANDOFF §8 #42
table(doc, ["", "調べること", "1人あたり", "アンケート"],
      [["**UI研究**", "どの入れ方だと答えやすいか", "**2回**", "毎回2つ。桃山台駅 6問、南口 16問（2回目は19問）"],
       ["**計数研究**", "人数をどう数えるのが正確か", "**1回**", "南口で1つ（5問）"]],
      widths=[24, 56, 20, 56])
p(doc, "別の研究として数える。同じ人が両方に出てもよい。")

# ---------------------------------------------------------------- UI研究
h2(doc, "UI研究")
p(doc, "参加者は **Aさん〜Lさん**（予備 Mさん〜Xさん）。1回＝ＪＲ吹田駅（南口）からの**往復2本**。")

h3(doc, "入力方式")
table(doc, ["", "ゆき", "かえり"],
      [["**1回目**", "**テンキー**", "**＋−ボタン**"],
       ["**2回目**", "**＋−ボタン**", "**テンキー**"]],
      widths=[22, 44, 44])

# ★2026-10-07: 全員「なし→あり」だと、ボタンの効果と2回目の慣れが区別できない。半分の人は順番を逆にする
h3(doc, "わからないボタン（往復とも同じ）")
table(doc, ["", "ID", "予備", "1回目", "2回目"],
      [["**奇数組**", "A C E G I K", "M O Q S U W", "なし", "**あり**"],
       ["**偶数組**", "B D F H J L", "N P R T V X", "**あり**", "なし"]],
      widths=[22, 38, 38, 24, 24])
p(doc, "条件はアプリが **QR・ゆき／かえり・ID** から決める。")
p(doc, "**予備IDは、抜けた人と同じ組から使う。**")

h3(doc, "乗ってもらい方")
steps(doc, [
    "1回目と2回目は、**昼と夕方で1回ずつ**乗ってもらう",
    "日付は別の日でよい。**順番も自由**（夕方が先でもよい）",
    "都合がつかなければ同じ時間帯が2回でもよい",
    "**同じバスに1回目の人と2回目の人が混ざってよい**",
    "アンケートは2つ。桃山台駅でゆきの分、南口でかえりの分を、配布用紙の裏のQRから各自で答える",
])

h3(doc, "使う便")
table(doc, ["", "集合", "ゆき 南口発", "桃山台着", "かえり 桃山台発", "南口着・解散"],
      [["**昼**", "11:35", "**12:03**", "12:48", "**13:48**", "14:32"],
       ["**夕方**", "15:25", "**15:55**", "16:40", "**17:18**", "18:03"]],
      widths=[18, 20, 34, 24, 40, 32])
# ★2026-10-07: かえりは松本がいないので、桃山台側の乗り方も書く（参加者向け案内にも同じことを書いた）
p(doc, "すべて阪急バス 吹田市内線2系統。南口のりば2からは **[2]** と **[3]** の両方が"
       "「桃山台駅ゆき」で出る。**乗るのは [2]。**"
       "かえりは桃山台駅 **2番のりば**の **[2]「ＪＲ吹田駅（南口）」**。**終点まで乗る。**")

# ★2026-10-07: ゆき/かえり別のQRをやめて「回」ごとにした。桃山台でQRを出す人がいないため。
# 計数研究のQRは ★乗降人数計算 に分けたので2枚
# ★2026-10-08: 当日の説明と桃山台での迷いを減らすため、時間割・手順・アプリとアンケートのQRを1枚にまとめた
h3(doc, "配布用紙（2種類）")
table(doc, ["配布用紙", "渡す人"],
      [["**UI研究 1回目**", "UI研究で1回目の人"],
       ["**UI研究 2回目**", "UI研究で2回目の人"]],
      widths=[34, 96])
p(doc, "**名簿でその人が何回目かを確かめてから渡す。**参加者に条件は言わない。QRポスターは予備。")

h3(doc, "南口ですること")
# ★2026-10-07: 位置情報は報告画面が出た時点で聞かれるので、練習の前に置く
steps(doc, [
    "名簿で何回目かを確かめ、その回の配布用紙を渡す。1回目の人にはカードも渡す",
    "配布用紙の表のQRを読んでもらう",
    "ログインしたら「**ゆき　吹田駅 → 桃山台駅**」を選んでもらう。画面の「**○回目　ゆき**」を確かめ、"
    "位置情報は「**許可**」してもらう",
    # ★2026-10-08: iPhone の Safari で許可を聞かれず、オレンジの帯のままだったことがある（HANDOFF §8 #43）
    "画面の下の帯が「**位置情報 OK**」になっているか確かめる。オレンジの帯なら帯を押し、出てきた手順を一緒にやる",
    "のりばで練習。「**練習**」を押す → のりばで待っている人の数を**2回**答える → "
    "画面の上の「**練習モード**」をタップして終える",
    "乗るバスを指定して送り出す。"
    "「**桃山台駅で降りたら『かえりに切り替える』を押してください**」と一言伝える",
    "送り出したら離脱する。**松本はバスに乗らない**",
])

h3(doc, "かえりの切り替え")
table(doc, None,
      [["**QR**", "南口で1回読むだけ"],
       ["**切り替え**", "桃山台駅で参加者がアプリの「**かえりに切り替える**」を押す"],
       ["**押し忘れ**", "ゆきを選んで70分たつと、画面に押すよう案内が出る"],
       ["**押し間違い**", "「**ゆきに戻す**」で戻せる"]],
      widths=[28, 128])

h3(doc, "数え方（参加者に伝える）")
table(doc, None,
      [["**数に入れる人**", "運転士以外で、車内にいる全員。**自分も入れる**"],
       ["**いつ**", "停まったバス停を**発車したら**数えて答え、**次のバス停までに送る**"]],
      widths=[28, 128])

# ---------------------------------------------------------------- 計数研究
# ★2026-10-07: 計数研究は別の実験として資料を分けた（案内・QR・アンケートは ../★乗降人数計算/）
h2(doc, "計数研究")
p(doc, "計数研究は **★乗降人数計算** の資料で行う。")
# ★2026-10-09: 位置情報の帯は人数アプリ・乗降アプリの両方にあるので、計数研究でも南口で確かめる
p(doc, "南口では UI研究と同じく、画面の下の帯が「**位置情報 OK**」になっているか確かめる。")

# ---------------------------------------------------------------- 持ちもの
h2(doc, "持っていくもの")
table(doc, None,
      [["配布カード", "Aさん〜Lさん、予備 Mさん〜Xさん。ID と パスワード"],
       ["配布用紙", "1回目用・2回目用を人数分（★UI研究）"],
       ["QRポスター", "予備。2枚（1回目／2回目）"],
       ["名簿", "その人が何回目か"],
       ["対応表", "accounts.csv。人目に触れない場所に"]],
      widths=[32, 124])

doc.save(OUT)
print("書き出しました: " + OUT)
